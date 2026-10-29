"""Editable OOXML and direct PDF indices; never convert user documents here."""
from io import BytesIO
from urllib.parse import quote, unquote
from zipfile import ZipFile
from xml.etree import ElementTree as E

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject, FloatObject, NumberObject, BooleanObject
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from . import pdf

LABELS = {'ru':{'identifier':'Обозначение','title':'Название','pages':'Страниц','path':'Файл'},
          'en':{'identifier':'Exhibit','title':'Title','pages':'Pages','path':'File'}}


def columns(config):
    return ['identifier','title']+[k for k in ('pages','path') if k in config['columns']]


def check_text(config, rows):
    pdf.init_font()
    glyphs = pdfmetrics.getFont('DejaVu').face.charToGlyph
    for value in [config['title'],*LABELS[config['language']].values(),*[str(r[k]) for r in rows for k in columns(config)]]:
        for char in value:
            if char != '\n' and (ord(char) not in glyphs or ord(char) < 32):
                raise ValueError(f'Шрифт перечня не поддерживает символ U+{ord(char):04X}. Измените текст перечня или название приложения.')


def docx(config, rows):
    doc = Document();section = doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.left_margin=section.right_margin=Mm(18)
    section.top_margin=section.bottom_margin=Mm(18)
    normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(0)
    doc.styles['Heading 1'].font.color.rgb=RGBColor.from_string('122A4B')
    doc.core_properties.author='BL Exhibit Manager';doc.core_properties.title=config['title']
    doc.add_heading(config['title'],level=1)
    cols=columns(config);table=doc.add_table(rows=1,cols=len(cols));table.style='Table Grid'
    table.autofit=False
    widths=column_widths(config)
    for column,key in zip(table.columns,cols):column.width=Pt(widths[key])
    header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
    for i,key in enumerate(cols):table.rows[0].cells[i].text=LABELS[config['language']][key]
    for row in rows:
        cells=table.add_row().cells
        for i,key in enumerate(cols):
            if key != 'title':cells[i].text=str(row[key]);continue
            paragraph=cells[i].paragraphs[0]
            rid=doc.part.relate_to(quote(row['path'],safe='/'),RT.HYPERLINK,is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rid)
            run=OxmlElement('w:r');props=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'122A4B');props.append(color);run.append(props)
            text=OxmlElement('w:t');text.set(qn('xml:space'),'preserve');text.text=row['title'];run.append(text)
            link.append(run);paragraph._p.append(link)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,key in zip(row.cells,cols):
            cell.width=Pt(widths[key]);margins=OxmlElement('w:tcMar')
            for side in ('top','bottom'):
                item=OxmlElement('w:'+side);item.set(qn('w:w'),'60');item.set(qn('w:type'),'dxa');margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:run.bold=True;run.font.size=Pt(9)
    paragraph=section.footer.paragraphs[0]
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');paragraph._p.append(field)
    out=BytesIO();doc.save(out)
    # Package timestamps are not a user-visible index change.
    stable=BytesIO()
    with ZipFile(BytesIO(out.getvalue())) as source, ZipFile(stable,'w') as target:
        for info in source.infolist():
            info.date_time=(1980,1,1,0,0,0);target.writestr(info,source.read(info))
    return stable.getvalue()


def wrap(value, width, size=9):
    lines=[]
    for block in str(value).split('\n'):
        line=''
        for char in block:
            if line and pdfmetrics.stringWidth(line+char,'DejaVu',size)>width:
                lines.append(line);line=''
            line+=char
        lines.append(line)
    return lines


def column_widths(config):
    widths={'identifier':92,'pages':58,'path':150}
    widths['title']=A4[0]-102-sum(widths[k] for k in columns(config) if k!='title')
    return widths


def check_layout(config, rows):
    widths=column_widths(config)
    available=A4[1]-55-16*len(wrap(config['title'],A4[0]-102,12))-26-52
    for row in rows:
        if 13*max(len(wrap(row[k],widths[k]-8)) for k in columns(config))+10>available:
            raise ValueError('Строка перечня не помещается на страницу. Сократите название или путь приложения.')


def pdf_document(config, rows):
    out=BytesIO();c=canvas.Canvas(out,pagesize=A4,invariant=1);c.setTitle(config['title']);c.setAuthor('BL Exhibit Manager')
    page_width,page_height=A4;left=51;usable=page_width-2*left
    cols=columns(config)
    widths=column_widths(config)
    boxes=[];page=-1;y=0
    def header():
        nonlocal page,y
        page+=1;c.setFont('DejaVu',12);y=page_height-55
        for line in wrap(config['title'],usable,12):c.drawString(left,y,line);y-=16
        y-=12;x=left;c.setFont('DejaVu',9)
        for key in cols:c.drawString(x+4,y,LABELS[config['language']][key]);x+=widths[key]
        y-=10;c.line(left,y,left+usable,y);y-=4
        c.setFont('DejaVu',9);c.drawCentredString(page_width/2,27,str(page+1))
    header()
    body_top=y
    for row in rows:
        lines={k:wrap(row[k],widths[k]-8) for k in cols}
        height=13*max(len(v) for v in lines.values())+10
        if height>body_top-52:raise ValueError('Строка перечня не помещается на страницу. Сократите название или путь приложения.')
        if y-height<52:c.showPage();header()
        x=left
        for key in cols:
            c.setFont('DejaVu',9);c.setFillColorRGB(*((0.07,0.16,0.29) if key=='title' else (0,0,0)))
            for i,line in enumerate(lines[key]):
                baseline=y-13*(i+1)
                c.drawString(x+4,baseline,line)
                if key=='title' and line:
                    boxes.append((page,row['path'],[x+4,baseline-2,x+4+pdfmetrics.stringWidth(line,'DejaVu',9),baseline+9]))
            x+=widths[key]
        y-=height;c.setStrokeColorRGB(.85,.88,.93);c.line(left,y,left+usable,y)
    c.save();writer=PdfWriter();writer.append(PdfReader(BytesIO(out.getvalue())))
    for page,path,rect in boxes:
        spec=DictionaryObject({NameObject('/Type'):NameObject('/Filespec'),NameObject('/F'):TextStringObject(path),NameObject('/UF'):TextStringObject(path)})
        action=DictionaryObject({NameObject('/S'):NameObject('/GoToR'),NameObject('/F'):spec,
                                 NameObject('/D'):ArrayObject([NumberObject(0),NameObject('/Fit')]),NameObject('/NewWindow'):BooleanObject(True)})
        annotation=DictionaryObject({NameObject('/Type'):NameObject('/Annot'),NameObject('/Subtype'):NameObject('/Link'),
            NameObject('/Rect'):ArrayObject([FloatObject(x) for x in rect]),NameObject('/Border'):ArrayObject([NumberObject(0)]*3),NameObject('/A'):action})
        writer.add_annotation(page,annotation)
    result=BytesIO();writer.write(result);return result.getvalue()


def generate(fmt, config, rows):
    check_text(config,rows)
    check_layout(config,rows)
    if fmt not in ('docx','pdf'):raise ValueError('Неизвестный формат перечня.')
    return docx(config,rows) if fmt=='docx' else pdf_document(config,rows)


def targets(fmt, data):
    if fmt=='docx':
        with ZipFile(BytesIO(data)) as z:
            rels=E.fromstring(z.read('word/_rels/document.xml.rels'))
            return [unquote(r.attrib['Target']) for r in rels if r.attrib.get('Type','').endswith('/hyperlink')]
    result=[]
    for page in PdfReader(BytesIO(data)).pages:
        for ref in page.get('/Annots',[]):
            action=ref.get_object().get('/A',{})
            if action.get('/S')!='/GoToR':raise ValueError('Непереносимая ссылка перечня.')
            spec=action['/F'].get_object();result.append(str(spec['/UF']))
    return result


def verify_package(data, config, rows):
    expected=[r['path'] for r in rows]
    if len(set(expected))!=len(rows):raise ValueError('Повторная строка перечня.')
    with ZipFile(BytesIO(data)) as archive:
        for row in rows:
            name='Submission/'+row['path']
            if name not in archive.namelist() or len(PdfReader(BytesIO(archive.read(name))).pages)!=row['pages']:
                raise ValueError('Перечень не соответствует выданным PDF.')
        for fmt in config['formats']:
            value=archive.read('Submission/Exhibit index.'+fmt)
            found=targets(fmt,value)
            if set(found)!=set(expected) or any(p not in expected for p in found):raise ValueError('Неверная цель ссылки перечня.')
            if fmt=='docx':
                with ZipFile(BytesIO(value)) as doc:
                    rels={r.attrib['Id']:r.attrib for r in E.fromstring(doc.read('word/_rels/document.xml.rels'))}
                    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
                    actual=E.fromstring(doc.read('word/document.xml')).findall('.//w:tbl/w:tr',ns)[1:]
                    if len(actual)!=len(rows):raise ValueError('Неполная таблица перечня.')
                    for tr,row in zip(actual,rows):
                        links=tr.findall('.//w:hyperlink',ns)
                        if len(links)!=1:raise ValueError('Неверные ссылки строки перечня.')
                        link=links[0];rel=rels[link.attrib['{'+ns['r']+'}id']]
                        if rel.get('TargetMode')!='External' or unquote(rel['Target'])!=row['path'] or ''.join(t.text or '' for t in link.findall('.//w:t',ns))!=row['title']:
                            raise ValueError('Название и цель строки перечня не совпадают.')
