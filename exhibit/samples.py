"""Reproducible fictional fixtures. No material from actual proceedings."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import argparse

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from lxml import etree as E
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

from .pdf import init_font
from .word import W, R, REL, NS, FOOT, RELS, xml
from .project import Store


def make_pdf(title, pages):
    init_font()
    out = BytesIO()
    c = canvas.Canvas(out, pagesize=A4, invariant=1)
    c.setTitle("FICTIONAL TRAINING MATERIAL")
    for number, lines in enumerate(pages, 1):
        c.setFont("DejaVu", 9)
        c.setFillColorRGB(.35, .35, .35)
        c.drawString(52, 796, "УЧЕБНЫЙ МАТЕРИАЛ / FICTIONAL TRAINING MATERIAL")
        c.setFillColorRGB(0, 0, 0)
        c.setFont("DejaVu", 14)
        c.drawString(52, 751, title)
        c.setFont("DejaVu", 11)
        y = 704
        for line in lines:
            c.drawString(52, y, line)
            y -= 27
        c.setFont("DejaVu", 9)
        c.drawString(52, 46, "Все организации, события и утверждения вымышлены.")
        c.drawRightString(543, 46, str(number))
        c.showPage()
    c.save()
    return out.getvalue()


def main_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.8)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(11)
    style.paragraph_format.space_after = Pt(8)
    for name in ("Title", "Heading 1"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        for border in doc.styles[name]._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.core_properties.author = "Buzko Legal training fixture"
    doc.core_properties.title = "Учебная подача"
    doc.add_paragraph("Учебная подача", "Title")
    doc.add_paragraph("Вымышленный спор между «Северный маяк» и «Тихая бухта». Документ создан исключительно для проверки обработки приложений и сносок; он не предназначен для подачи в суд.")
    doc.add_paragraph("Позиция учебной стороны", "Heading 1")
    statements = ["1. Стороны обсудили сроки передачи учебного оборудования.",
                  "2. Правило повторно упоминается в следующем доводе.",
                  "3. Два материала необходимо рассматривать совместно.",
                  "4. Требования изложены в отдельном документе.",
                  "5. Дополнительный материал намеренно не включён в загрузку."]
    for i, sentence in enumerate(statements, 1):
        p = doc.add_paragraph(sentence)
        run = p.add_run()
        prop = E.SubElement(run._r, f"{{{W}}}rPr")
        E.SubElement(prop, f"{{{W}}}vertAlign", {f"{{{W}}}val": "superscript"})
        E.SubElement(run._r, f"{{{W}}}footnoteReference", {f"{{{W}}}id": str(i)})
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text, table.cell(0, 1).text = "Учебная дата", "Событие"
    table.cell(1, 0).text, table.cell(1, 1).text = "14 марта 2026", "Передача макета оборудования"
    sec.header.paragraphs[0].text = "ТЕСТОВЫЙ ДОКУМЕНТ • НЕ ДЛЯ ПОДАЧИ"
    sec.footer.paragraphs[0].text = "Вымышленные материалы"
    raw = BytesIO()
    doc.save(raw)
    with ZipFile(raw) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    footnotes = E.Element(f"{{{W}}}footnotes", nsmap={"w": W, "r": R})
    for id_, kind, mark in [(-1, "separator", "separator"), (0, "continuationSeparator", "continuationSeparator")]:
        f = E.SubElement(footnotes, f"{{{W}}}footnote", {f"{{{W}}}id": str(id_), f"{{{W}}}type": kind})
        E.SubElement(E.SubElement(E.SubElement(f, f"{{{W}}}p"), f"{{{W}}}r"), f"{{{W}}}{mark}")
    texts = ["См. RLA-31, пункт 1.", "Повторная ссылка: RLA-31.", "См. RLA-31 и RLA-32, страницы 1 и 3.",
             "См. Statement of Claim, раздел 2.", "См. RLA-99 (намеренно отсутствует)."]
    for i, text in enumerate(texts, 1):
        f = E.SubElement(footnotes, f"{{{W}}}footnote", {f"{{{W}}}id": str(i)})
        p = E.SubElement(f, f"{{{W}}}p")
        pp = E.SubElement(p, f"{{{W}}}pPr")
        E.SubElement(pp, f"{{{W}}}pStyle", {f"{{{W}}}val": "FootnoteText"})
        run = E.SubElement(p, f"{{{W}}}r")
        rp = E.SubElement(run, f"{{{W}}}rPr")
        E.SubElement(rp, f"{{{W}}}vertAlign", {f"{{{W}}}val": "superscript"})
        E.SubElement(run, f"{{{W}}}footnoteRef")
        # Deliberately split identifiers across runs and preserve bold/italic.
        cut = text.find("RLA-") + 4 if "RLA-" in text else len(text)//2
        for j, chunk in enumerate([" " + text[:cut], text[cut:]]):
            run = E.SubElement(p, f"{{{W}}}r")
            rp = E.SubElement(run, f"{{{W}}}rPr")
            E.SubElement(rp, f"{{{W}}}sz", {f"{{{W}}}val": "20"})
            if i == 1:
                E.SubElement(rp, f"{{{W}}}{'b' if j == 0 else 'i'}")
            t = E.SubElement(run, f"{{{W}}}t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
            t.text = chunk
    parts[FOOT] = E.tostring(footnotes, xml_declaration=True, encoding="UTF-8", standalone=True)
    relationships = xml(parts["word/_rels/document.xml.rels"])
    E.SubElement(relationships, f"{{{REL}}}Relationship", Id="rIdTrainingFootnotes", Type=R+"/footnotes", Target="footnotes.xml")
    parts["word/_rels/document.xml.rels"] = E.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True)
    ct = xml(parts["[Content_Types].xml"])
    E.SubElement(ct, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", PartName="/word/footnotes.xml", ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml")
    parts["[Content_Types].xml"] = E.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as z:
        for name, value in parts.items():
            z.writestr(name, value)
    return out.getvalue()


def files():
    return {
        "RLA-31-original.pdf": make_pdf("Решение по учебному спору", [
            ["1. «Северный маяк» передал макет 14 марта 2026 года.", "2. «Тихая бухта» подтвердила получение макета.", "3. Указанные обстоятельства вымышлены."],
            ["4. Срок проверки составляет десять учебных дней.", "5. Документ не содержит действующих правовых норм."]]),
        "RLA-31-translation.pdf": make_pdf("Decision in a fictional dispute", [[
            "1. Northern Beacon delivered the model on 14 March 2026.", "2. Quiet Bay confirmed receipt of the model.",
            "3. These circumstances are fictional.", "4. The inspection period is ten training days.", "5. This document contains no applicable legal rules."]]),
        "RLA-32.pdf": make_pdf("Переписка учебных сторон", [
            ["1. Сообщение от 10 марта 2026 года.", "Просим подготовить учебный макет к передаче."],
            ["EXCLUDED_PAGE_SECRET_7E93", "Эта страница не должна попасть в итоговый PDF."],
            ["3. Сообщение от 14 марта 2026 года.", "INCLUDED_FRAGMENT_42", "Макет получен. Начинается учебная проверка.",
             "", "", "", "EXCLUDED_REGION_SECRET_5B71", "Эта нижняя часть страницы не должна сохраняться."]]),
        "R-1.pdf": make_pdf("Exhibit R-1 — готовый учебный файл", [["Сохраните этот файл без повторной обработки.", "Готовое оформление является частью исходного PDF."]]),
        "Statement of Claim.pdf": make_pdf("Statement of Claim", [["1. This claim is fictional and intended for training only.", "2. The claimant requests delivery of a model vessel."]]),
        "Main document.docx": main_docx(),
        "RLA-99.pdf": make_pdf("Дополнительный учебный материал RLA-99", [["Отдельный файл для устранения ненайденной ссылки."]]),
    }


def demo_project(store):
    p, content = store.create("Учебная подача"), files()
    for name, title, prefix, number, folder, mode in [
        ("RLA-31-original.pdf", "Решение по учебному спору", "RLA", 31, "RLA", "prepare"),
        ("RLA-32.pdf", "Переписка учебных сторон", "RLA", 32, "RLA", "prepare"),
        ("R-1.pdf", "Готовое приложение", "R", 1, "R", "passthrough"),
        ("Statement of Claim.pdf", "Statement of Claim", "", None, "Other documents", "passthrough")]:
        store.upload(p, name, content[name])
        d = p["documents"][-1]
        store.update(p, d["id"], {"title": title, "prefix": prefix, "number": number, "folder": folder, "mode": mode,
            "filename": f"{prefix}-{number}.pdf" if number else name, "designation": "Exhibit" if number else "",
            "language": "Русский" if number else "English", "original_label": "[Original in Russian]"})
    first, second = p["documents"][:2]
    store.upload(p, "RLA-31-translation.pdf", content["RLA-31-translation.pdf"], "translation", first["id"])
    store.update(p, first["id"], {"translation_label": "[Translation from Russian]"})
    store.update(p, second["id"], {"selection": [{"page": 1}, {"page": 3, "rect": [.07, .13, .94, .28], "omission_before": True, "omission_after": True}]})
    store.upload(p, "Main document.docx", content["Main document.docx"], "main")
    store.scan(p)
    return p


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="output/samples")
    parser.add_argument("--stress", type=int, default=0)
    args = parser.parse_args()
    folder = Path(args.output)
    folder.mkdir(parents=True, exist_ok=True)
    for name, data in files().items():
        (folder / name).write_bytes(data)
    for i in range(1, args.stress+1):
        (folder / f"LOAD-{i:03}.pdf").write_bytes(make_pdf(f"Учебный документ {i}", [[f"Вымышленный текст {i} для проверки нагрузки."]]))
    print(f"Created {7+args.stress} fictional files in {folder.resolve()}")
