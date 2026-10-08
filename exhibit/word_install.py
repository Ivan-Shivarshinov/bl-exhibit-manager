"""Opt-in local pilot sideload, following Office-Addin-Scripts registration.

No debugger, network share, Office policy edits, or closing Word.
"""
from io import BytesIO
import base64
import hashlib
import json
import os
from pathlib import Path
import plistlib
import ssl
import subprocess
import sys
import threading
import uuid
from urllib.parse import urlsplit
from zipfile import ZipFile, ZIP_DEFLATED

from . import word_setup
from .external_process import popen

ADDIN_ID = '43a6f2bd-d10f-4c7a-91b5-f3938c499ac4'
REGISTRY = r'Software\Microsoft\Office\16.0\WEF\Developer'


def registration_token(snapshot):
    return hashlib.sha256(json.dumps(snapshot, sort_keys=True).encode()).hexdigest()


def local_manifest(data):
    """Recognize our local add-in, not arbitrary content occupying its ID."""
    from lxml import etree
    try:
        if len(data) > 1024 * 1024: return False
        root = etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True))
        source = urlsplit(root.find('.//{*}SourceLocation').get('DefaultValue'))
        return (root.findtext('{*}Id', '').lower() == ADDIN_ID
                and root.find('{*}DisplayName').get('DefaultValue') == 'BL Exhibit Manager'
                and source.scheme == 'https' and source.hostname in ('localhost', '127.0.0.1', '::1')
                and source.path == '/word/index.html' and not source.username and not source.password)
    except (ValueError, AttributeError, etree.XMLSyntaxError):
        return False


class Cancelled(Exception):
    pass


def starter_document(root):
    """A new synthetic document each time; never overwrite a user's document."""
    from docx import Document
    from lxml import etree
    doc = Document()
    doc.add_heading('Проверка панели BL Exhibit Manager', 0)
    doc.add_paragraph('Это учебный документ. Ваши рабочие документы не изменены. Выберите учебную подачу в панели и создайте сноску. После проверки этот документ можно закрыть.')
    stream = BytesIO(); doc.save(stream)
    with ZipFile(BytesIO(stream.getvalue())) as source:
        files = {name: source.read(name) for name in source.namelist()}
    ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    rels = etree.fromstring(files['_rels/.rels'])
    etree.SubElement(rels, f'{{{ns}}}Relationship', Id='rIdBLPane',
        Type='http://schemas.microsoft.com/office/2011/relationships/webextensiontaskpanes', Target='word/webextensions/taskpanes.xml')
    files['_rels/.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8')
    types_ns = 'http://schemas.openxmlformats.org/package/2006/content-types'
    types = etree.fromstring(files['[Content_Types].xml'])
    for part, content in [('taskpanes', 'webextensiontaskpanes'), ('webextension', 'webextension')]:
        etree.SubElement(types, f'{{{types_ns}}}Override', PartName=f'/word/webextensions/{part}.xml', ContentType=f'application/vnd.ms-office.{content}+xml')
    files['[Content_Types].xml'] = etree.tostring(types, xml_declaration=True, encoding='UTF-8')
    files['word/webextensions/taskpanes.xml'] = b'<wetp:taskpanes xmlns:wetp="http://schemas.microsoft.com/office/webextensions/taskpanes/2010/11" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><wetp:taskpane dockstate="right" visibility="1" width="350" row="1"><wetp:webextensionref r:id="rIdBLPane"/></wetp:taskpane></wetp:taskpanes>'
    files['word/webextensions/_rels/taskpanes.xml.rels'] = f'<Relationships xmlns="{ns}"><Relationship Id="rIdBLPane" Type="http://schemas.microsoft.com/office/2011/relationships/webextension" Target="webextension.xml"/></Relationships>'.encode()
    files['word/webextensions/webextension.xml'] = f'<we:webextension xmlns:we="http://schemas.microsoft.com/office/webextensions/webextension/2010/11" id="{{{ADDIN_ID}}}"><we:reference id="{ADDIN_ID}" version="0.2.0.0" store="developer" storeType="Registry"/><we:alternateReferences/><we:properties/><we:bindings/></we:webextension>'.encode()
    target = root / 'word-local' / 'starter'
    target.mkdir(parents=True, exist_ok=True)
    path = target / f'Word-check-{uuid.uuid4().hex}.docx'
    stream = BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as result:
        for name, content in files.items(): result.writestr(name, content)
    word_setup.atomic_write(path, stream.getvalue())
    return path


class NativeSettings:
    def __init__(self, platform=None, home=None, registry=REGISTRY, keychain=None, admin_trust=False):
        self.platform = platform or sys.platform
        self.home = Path(home) if home else Path.home()
        self.registry = registry
        self.keychain = Path(keychain) if keychain else self.home / 'Library/Keychains/login.keychain-db'
        self.admin_trust = admin_trust  # Only used by disposable CI verification.

    def run(self, args, cancel=None, timeout=120, check=True):
        flags = {'creationflags': subprocess.CREATE_NO_WINDOW} if sys.platform == 'win32' else {}
        process = popen(args, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE, **flags)
        import time
        deadline = time.monotonic() + timeout
        try:
            while True:
                if cancel and cancel.is_set(): raise Cancelled()
                if time.monotonic() >= deadline: raise ValueError('Системная операция не завершилась вовремя. Повторите действие; проекты и документы сохранены.')
                try:
                    output, _ = process.communicate(timeout=.25)
                    if check and process.returncode: raise ValueError('Система не разрешила настройку. Можно повторить подключение или продолжить работу без панели.')
                    return process.returncode, output
                except subprocess.TimeoutExpired: pass
        finally:
            if process.poll() is None:
                process.terminate()
                try: process.communicate(timeout=3)
                except subprocess.TimeoutExpired: process.kill(); process.communicate()

    def word(self):
        if self.platform == 'win32':
            import winreg
            for hive in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
                for view in (winreg.KEY_WOW64_64KEY, winreg.KEY_WOW64_32KEY):
                    try:
                        with winreg.OpenKey(hive, r'SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE', 0, winreg.KEY_READ | view) as key:
                            path = Path(winreg.QueryValueEx(key, '')[0])
                        if path.is_file(): return path
                    except FileNotFoundError: pass
            raise ValueError('Microsoft Word не найден. Установите настольный Word с поддержкой WordApi 1.5; обычная работа доступна без него.')
        if self.platform == 'darwin':
            for path in (Path('/Applications/Microsoft Word.app'), self.home / 'Applications/Microsoft Word.app'):
                if (path / 'Contents/Info.plist').is_file(): return path
            raise ValueError('Microsoft Word не найден. Установите настольный Word с поддержкой WordApi 1.5; обычная работа доступна без него.')
        raise ValueError('Экспериментальное подключение поддерживается на Windows и macOS.')

    def check_word(self):
        path = self.word()
        if self.platform == 'darwin':
            with (path / 'Contents/Info.plist').open('rb') as file: version = plistlib.load(file).get('CFBundleShortVersionString', '')
            parts = tuple(int(x) for x in version.split('.') if x.isdigit())
            supported = parts >= (16, 70)
        else:
            import ctypes
            from ctypes import wintypes
            dll = ctypes.WinDLL('version', use_last_error=True)
            dll.GetFileVersionInfoSizeW.argtypes = [wintypes.LPCWSTR, ctypes.POINTER(wintypes.DWORD)]
            dll.GetFileVersionInfoW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p]
            dll.VerQueryValueW.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR, ctypes.POINTER(ctypes.c_void_p), ctypes.POINTER(wintypes.UINT)]
            size = dll.GetFileVersionInfoSizeW(str(path), None)
            buffer = ctypes.create_string_buffer(size)
            pointer, length = ctypes.c_void_p(), wintypes.UINT()
            if not size or not dll.GetFileVersionInfoW(str(path), 0, size, buffer) or not dll.VerQueryValueW(buffer, '\\', ctypes.byref(pointer), ctypes.byref(length)):
                raise ValueError('Не удалось определить версию Word. Проверьте установку Word и повторите настройку.')
            info = ctypes.cast(pointer, ctypes.POINTER(wintypes.DWORD))
            parts = (info[2] >> 16, info[2] & 65535, info[3] >> 16, info[3] & 65535)
            supported = parts >= (16, 0, 16130, 20332)
        if not supported:
            raise ValueError('Этой версии Word недоступны нужные функции панели. Нужен WordApi 1.5: совместимая версия Microsoft 365 или Office 2024. Проверьте обновления вашей редакции Word.')
        return path

    def _mac_target(self):
        folder = self.home / 'Library/Containers/com.microsoft.Word/Data/Documents/wef'
        if not folder.resolve().is_relative_to(self.home.resolve()):
            raise ValueError('Каталог Word находится вне профиля пользователя. Автоматическая настройка недоступна.')
        return folder / f'{ADDIN_ID}.BLExhibitManager.Word.xml'

    def registration_snapshot(self):
        if self.platform == 'win32':
            import winreg
            try:
                with winreg.OpenKey(winreg.HKEY_CURRENT_USER, self.registry) as key:
                    value, kind = winreg.QueryValueEx(key, ADDIN_ID)
                    if kind != winreg.REG_SZ: raise ValueError('Существующая регистрация панели имеет другой формат. Она сохранена.')
                return {'kind': 'windows', 'value': value}
            except FileNotFoundError: return None
        target = self._mac_target()
        if target.is_symlink(): raise ValueError('Регистрация панели является ссылкой на другой файл. Она сохранена; автоматическая замена недоступна.')
        if target.exists():
            if target.stat().st_size > 1024 * 1024: raise ValueError('Файл регистрации панели не распознан. Он сохранён.')
            return {'kind': 'mac', 'value': base64.b64encode(target.read_bytes()).decode(), 'mode': target.stat().st_mode & 0o777}
        return None

    def registration_status(self, manifest):
        snapshot = self.registration_snapshot()
        if snapshot is None: return {'state': 'missing'}
        if snapshot['kind'] == 'windows':
            same = os.path.normcase(snapshot['value']) == os.path.normcase(str(manifest.resolve()))
            try:
                path = Path(snapshot['value'])
                data = path.read_bytes() if path.stat().st_size <= 1024 * 1024 else b''
            except OSError: data = b''
        else:
            data = base64.b64decode(snapshot['value'])
            same = manifest.is_file() and data == manifest.read_bytes()
        if same: return {'state': 'current'}
        if local_manifest(data): return {'state': 'previous', 'token': registration_token(snapshot)}
        return {'state': 'conflict'}

    def registration(self, manifest):
        state = self.registration_status(manifest)['state']
        if state not in ('current', 'missing'):
            raise ValueError('В Word уже есть другая регистрация панели. Вернитесь к настройке и подтвердите обновление, если оно доступно. Прежняя запись сохранена.')
        return state == 'current'

    def register(self, manifest, previous=None):
        if previous is not None:
            if self.registration_snapshot() != previous:
                raise ValueError('Регистрация Word изменилась после подтверждения. Она сохранена; откройте настройку заново.')
        elif self.registration(manifest): return
        if self.platform == 'win32':
            import winreg
            with winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, self.registry, 0, winreg.KEY_SET_VALUE) as key:
                winreg.SetValueEx(key, ADDIN_ID, 0, winreg.REG_SZ, str(manifest.resolve()))
        else:
            target = self._mac_target(); target.parent.mkdir(parents=True, exist_ok=True)
            word_setup.atomic_write(target, manifest.read_bytes(), 0o644)

    def restore_registration(self, manifest, previous=None):
        if previous is None:
            self.unregister(manifest); return
        snapshot = self.registration_snapshot()
        if snapshot == previous: return
        if snapshot is not None and not self.registration(manifest):
            raise ValueError('Регистрация Word изменилась; прежняя копия сохранена, чужая запись не заменена.')
        if self.platform == 'win32':
            import winreg
            with winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, self.registry, 0, winreg.KEY_SET_VALUE) as key:
                winreg.SetValueEx(key, ADDIN_ID, 0, winreg.REG_SZ, previous['value'])
        else:
            target = self._mac_target(); target.parent.mkdir(parents=True, exist_ok=True)
            word_setup.atomic_write(target, base64.b64decode(previous['value']), previous['mode'])

    def unregister(self, manifest):
        if not self.registration(manifest): return
        if self.platform == 'win32':
            import winreg
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, self.registry, 0, winreg.KEY_SET_VALUE) as key:
                winreg.DeleteValue(key, ADDIN_ID)
        else: self._mac_target().unlink()

    def trusted(self, cert_path, cancel=None):
        if self.platform == 'win32':
            der = word_setup.x509.load_pem_x509_certificate(cert_path.read_bytes()).public_bytes(word_setup.serialization.Encoding.DER)
            return any(value == der and encoding == 'x509_asn' and (trust is True or ssl.Purpose.SERVER_AUTH.oid in trust)
                       for value, encoding, trust in ssl.enum_certificates('ROOT'))
        code, _ = self.run(['/usr/bin/security', 'verify-cert', '-c', str(cert_path), '-p', 'ssl', '-s', 'localhost'], cancel, timeout=10, check=False)
        return code == 0

    def certificate_present(self, cert_path):
        cert = word_setup.x509.load_pem_x509_certificate(cert_path.read_bytes())
        if self.platform == 'win32':
            der = cert.public_bytes(word_setup.serialization.Encoding.DER)
            return any(value == der and encoding == 'x509_asn' for value, encoding, _ in ssl.enum_certificates('ROOT'))
        _, output = self.run(['/usr/bin/security', 'find-certificate', '-a', '-Z', str(self.keychain)], timeout=10, check=False)
        return cert.fingerprint(word_setup.hashes.SHA1()).hex().upper().encode() in output.upper()

    def install_trust(self, cert_path, cancel):
        if self.platform == 'win32':
            from tempfile import TemporaryDirectory
            from .launch import command
            with TemporaryDirectory(prefix='bl-word-trust-') as folder:
                result = Path(folder)/'result.json'
                args = command('--windows-root-add', str(cert_path), str(result))
                if self.admin_trust: args.append('--windows-root-machine')
                code, _ = self.run(args, cancel, check=False)
                try:
                    value = json.loads(result.read_text('utf-8'))
                    if not isinstance(value, dict): raise ValueError()
                except (OSError, ValueError):
                    raise ValueError('Не удалось получить результат добавления сертификата Windows. Повторите подключение; проекты сохранены.')
                if code or value.get('error'):
                    raise ValueError(value.get('error') or 'Windows не завершила добавление сертификата. Можно повторить подключение.')
                if type(value.get('added')) is not bool:
                    raise ValueError('Не удалось проверить результат добавления сертификата Windows. Повторите подключение.')
                return value['added']
        else:
            args = ['/usr/bin/security', 'add-trusted-cert', '-r', 'trustRoot', '-p', 'ssl', '-s', 'localhost', '-k', str(self.keychain), str(cert_path)]
            if self.admin_trust: args = ['/usr/bin/sudo', '-n', *args[:2], '-d', *args[2:]]
            self.run(args, cancel)

    def remove_trust(self, cert_path):
        if not self.trusted(cert_path): return
        if self.platform == 'win32':
            from . import windows_certificates
            certificate = word_setup.x509.load_pem_x509_certificate(cert_path.read_bytes())
            windows_certificates.remove(certificate.public_bytes(word_setup.serialization.Encoding.DER), admin=self.admin_trust)
        else:
            # Remove this certificate's SSL trust, retaining our reusable item.
            # RemoveTrustSettings can hang on current macOS. Replace only our
            # own SSL/localhost rule; never alter authorizationdb or other certs.
            args = ['/usr/bin/security', 'add-trusted-cert', '-r', 'deny', '-p', 'ssl', '-s', 'localhost', '-k', str(self.keychain), str(cert_path)]
            if self.admin_trust: args = ['/usr/bin/sudo', '-n', *args[:2], '-d', *args[2:]]
            self.run(args)
        if self.trusted(cert_path):
            raise ValueError('Система ещё доверяет сертификату приложения. Повторите удаление подключения и подтвердите системный запрос.')

    def open_word(self, path):
        word = self.check_word()
        if self.platform == 'win32': popen([str(word), '/n', str(path)])
        else: self.run(['/usr/bin/open', '-a', str(word), str(path)], timeout=15)


class Installer:
    def __init__(self, root, app, settings=None):
        self.root, self.app = Path(root), app
        self.settings = settings or NativeSettings()
        self.lock = threading.RLock()
        self.cancel = threading.Event()
        self.thread = None
        self.progress = None

    def _record(self):
        path = self.root / 'word-installation.json'
        if not path.exists(): return {}
        try:
            data = json.loads(path.read_text('utf-8'))
            fields = {'platform','fingerprint','trust_owned','certificate_owned','registration_owned','installed'}
            if not isinstance(data, dict) or set(data) - {'previous_registration'} != fields or data.get('platform') != self.settings.platform:
                raise ValueError()
            if not isinstance(data['fingerprint'], str) or len(data['fingerprint']) != 64 or any(x not in '0123456789abcdef' for x in data['fingerprint']): raise ValueError()
            if any(type(data[key]) is not bool for key in ('trust_owned','certificate_owned','registration_owned','installed')): raise ValueError()
            snapshot = data.get('previous_registration')
            if snapshot is not None:
                if not isinstance(snapshot, dict) or not isinstance(snapshot.get('value'), str) or len(snapshot['value']) > 1400000: raise ValueError()
                if self.settings.platform == 'win32':
                    if set(snapshot) != {'kind', 'value'} or snapshot['kind'] != 'windows': raise ValueError()
                elif self.settings.platform == 'darwin':
                    if set(snapshot) != {'kind', 'value', 'mode'} or snapshot['kind'] != 'mac' or type(snapshot['mode']) is not int or not 0 <= snapshot['mode'] <= 0o777: raise ValueError()
                    base64.b64decode(snapshot['value'], validate=True)
                else: raise ValueError()
            return data
        except (OSError, ValueError): raise ValueError('Не удалось прочитать состояние установки панели. Существующие настройки сохранены.')

    def _save(self, record):
        word_setup.atomic_write(self.root / 'word-installation.json', json.dumps(record).encode())

    def status(self):
        with self.lock:
            busy = bool(self.thread and self.thread.is_alive())
            try:
                record = self._record()
                registration = self.settings.registration_status(self.root / 'BLExhibitManager.Word.xml')
            except (OSError, ValueError) as exc:
                return {'phase':'error', 'message':str(exc) if isinstance(exc, ValueError) else 'Не удалось проверить регистрацию Word. Повторите настройку.', 'busy':busy}
            installed = bool(record.get('installed') and registration['state'] == 'current')
            extra = {'installed':installed, 'can_remove':bool(record.get('trust_owned') or record.get('registration_owned')), 'busy':busy}
            if registration['state'] == 'previous':
                extra.update(replacement_token=registration['token'], registration_message='Найдено прежнее подключение BL Exhibit Manager. Обновите его, чтобы панель работала с этим приложением.')
            elif registration['state'] == 'conflict':
                extra.update(registration_blocked=True, registration_message='Регистрация Word не распознана как наша локальная панель. Она сохранена. Обновление недоступно; скопируйте сводку диагностики для сообщения об ошибке. Обычная работа доступна без панели.')
            progress = self.progress
            path = self.root / 'word-installation-status.json'
            if progress is None and path.exists():
                try:
                    progress = json.loads(path.read_text('utf-8'))
                    if not isinstance(progress, dict) or set(progress) != {'phase','message'} or not all(isinstance(v, str) for v in progress.values()): raise ValueError()
                    if progress['phase'] not in ('error','cancelled','removed','waiting_word'):
                        progress = {'phase':'error', 'message':'Предыдущая настройка прервалась. Повторите подключение; сохранённые проекты доступны.'}
                except (OSError, ValueError):
                    progress = {'phase':'error', 'message':'Не удалось прочитать результат прежней настройки. Повторите подключение; проекты доступны.'}
            if progress and (busy or progress['phase'] in ('error','cancelled','removed')):
                return {**progress, **extra}
            if extra.get('registration_message'):
                return {**extra, 'phase':'needs_update' if registration['state']=='previous' else 'error', 'message':extra['registration_message']}
            if record.get('installed') and not installed:
                return {**extra, 'phase':'error', 'message':'Сохранённая регистрация Word отсутствует. Подключите панель заново; проекты сохранены.'}
            if progress: return {**progress, **extra}
            return {**extra, 'phase':'waiting_word' if installed else 'not_installed', 'message':'Настройка сохранена. Откройте панель в Word.' if installed else 'Панель подключается по вашему желанию.'}

    def _phase(self, phase, message):
        with self.lock:
            self.progress = {'phase':phase, 'message':message}
            self.root.mkdir(parents=True, exist_ok=True)
            word_setup.atomic_write(self.root / 'word-installation-status.json', json.dumps(self.progress).encode())

    def begin(self, remove=False, replacement_token=None):
        with self.lock:
            if self.thread and self.thread.is_alive(): return self.status()
            self.cancel.clear()
            self._phase('preparing', 'Проверяем Word и готовим подключение…')
            self.thread = threading.Thread(target=self._remove if remove else self._install, args=() if remove else (replacement_token,), name='word-local-setup', daemon=True)
            self.thread.start()
            return self.status()

    def stop(self):
        if not self.thread or not self.thread.is_alive(): return self.status()
        if self.progress and self.progress['phase'] == 'removing': return self.status()
        self.cancel.set()
        self._phase('cancelling', 'Отменяем настройку и сохраняем прежние данные…')
        return self.status()

    def _check_cancel(self):
        if self.cancel.is_set(): raise Cancelled()

    def _install(self, replacement_token=None):
        record = None; added_trust = added_registration = renewed = False
        try:
            self.settings.check_word(); self._check_cancel()
            manifest = self.root / 'BLExhibitManager.Word.xml'
            registration = self.settings.registration_status(manifest)
            replacement = None
            if registration['state'] == 'previous':
                replacement = self.settings.registration_snapshot()
                if not replacement_token or replacement_token != registration['token'] or replacement_token != registration_token(replacement):
                    raise ValueError('Найдено прежнее подключение панели. Нажмите «Обновить подключение» и подтвердите переход на это приложение.')
            elif registration['state'] == 'conflict':
                raise ValueError('Прежняя регистрация Word не распознана. Она сохранена; автоматическое обновление недоступно. Скопируйте сводку диагностики для сообщения об ошибке.')
            elif replacement_token:
                raise ValueError('Регистрация Word изменилась после подтверждения. Откройте настройку заново; существующая запись сохранена.')
            with word_setup.SETUP_LOCK:
                previous = self._record()
                word_setup.prepare(self.root, getattr(self.app.state, 'http_port', 8765))
                config, cert = word_setup.configuration(self.root, getattr(self.app.state, 'http_port', 8765))
                cert_path = Path(config['cert']); manifest = self.root / 'BLExhibitManager.Word.xml'
            registered = False if replacement else self.settings.registration(manifest)
            trusted = self.settings.trusted(cert_path, self.cancel)
            present = self.settings.certificate_present(cert_path)
            fingerprint = cert.fingerprint(word_setup.hashes.SHA256()).hex()
            if previous and previous.get('fingerprint') != fingerprint:
                raise ValueError('Сертификат изменился после установки. Прежнее доверие сохранено; автоматическая замена не выполнена.')
            if not trusted and not config.get('generated'):
                raise ValueError('Прежний сертификат не создан приложением. Его доверие сохранено; используйте настройку поставщика сертификата.')
            if present and not trusted and not previous.get('certificate_owned'):
                raise ValueError('Для этого сертификата уже есть другое системное правило доверия. Оно сохранено; автоматическая замена не выполнена.')
            record = {'platform':self.settings.platform, 'fingerprint':fingerprint,
                      'trust_owned': bool(previous.get('trust_owned') or not trusted),
                      'certificate_owned':bool(previous.get('certificate_owned') or not present),
                      'registration_owned': bool(previous.get('registration_owned') or not registered), 'installed':False,
                      'previous_registration': replacement or previous.get('previous_registration')}
            self._save(record)  # Ownership is persisted before any system mutation.
            self._check_cancel()
            if not trusted:
                self._phase('system_confirmation', 'Подтвердите запрос системы для локального сертификата, если он появился. Можно отменить настройку.')
                added_trust = True
                added = self.settings.install_trust(cert_path, self.cancel)
                self._check_cancel()
                trusted = self.settings.trusted(cert_path, self.cancel)
                if added is True and not trusted and self.settings.platform == 'win32':
                    # Windows may retain prior root approval after deletion while
                    # accepting a repeated import only into its temporary cache.
                    # Renew only our generated certificate, once, through the same
                    # protected importer. Never change ProtectedRoots or policies.
                    self._phase('preparing', 'Обновляем локальное подключение. Прежние сертификаты и документы сохранены…')
                    with word_setup.SETUP_LOCK:
                        old_config = (self.root / 'word-connection.json').read_bytes()
                        old_fingerprint = record['fingerprint']
                        try:
                            word_setup.prepare(self.root, getattr(self.app.state, 'http_port', 8765), renew=True)
                            config, cert = word_setup.configuration(self.root, getattr(self.app.state, 'http_port', 8765))
                            record['fingerprint'] = cert.fingerprint(word_setup.hashes.SHA256()).hex()
                            self._save(record)
                        except (OSError, ValueError):
                            word_setup.atomic_write(self.root / 'word-connection.json', old_config)
                            record['fingerprint'] = old_fingerprint
                            raise
                        renewed = True
                        cert_path = Path(config['cert'])
                    # Do not hold the setup lock while waiting for HTTP tasks in
                    # the old listener: an in-flight diagnostic may need it.
                    from . import word_tls
                    if not word_tls.restart_optional(self.app, self.root, getattr(self.app.state, 'http_port', 8765)):
                        raise ValueError('Не удалось обновить соединение панели. Повторите подключение; документы сохранены.')
                    self._check_cancel()
                    self._phase('system_confirmation', 'Подтвердите запрос системы для нового локального сертификата BL Exhibit Manager localhost. Можно отменить настройку.')
                    added = self.settings.install_trust(cert_path, self.cancel)
                    self._check_cancel()
                    trusted = self.settings.trusted(cert_path, self.cancel)
                if added is False:
                    added_trust = False
                    record['trust_owned'] = bool(previous.get('trust_owned')) and not renewed; self._save(record)
                if not trusted: raise ValueError('Система не подтвердила доверие сертификату приложения. Повторите подключение; обычная работа доступна.')
            self._check_cancel()
            self._phase('registering', 'Подключаем панель в настройках Word…')
            if not registered:
                added_registration = True; self.settings.register(manifest, previous=replacement)
            self._check_cancel()
            with word_setup.SETUP_LOCK:
                from . import word_tls
                listener = getattr(self.app.state, 'word_listener', None)
                if not listener or not listener[1].is_alive():
                    if not word_tls.start_optional(self.app, self.root, getattr(self.app.state, 'http_port', 8765)):
                        raise ValueError('Не удалось запустить соединение панели. Обычная работа доступна; повторите подключение.')
                verified = word_setup.status(self.root, self.app, verify=True)
                if verified['state'] != 'trusted': raise ValueError('Не удалось проверить защищённое соединение. Обычная работа доступна; повторите подключение.')
                self._check_cancel()
                record['installed'] = True; self._save(record)
            self._phase('opening_word', 'Открываем новый учебный документ в Word…')
            self._check_cancel()
            self.settings.open_word(starter_document(self.root))
            self._phase('waiting_word', 'Настройка выполнена. Дождитесь панели в Word и подтвердите её запуск, если Word попросит. Статус обновится автоматически.')
        except (Cancelled, ValueError, OSError) as exc:
            rollback_error = False
            if record and (not record.get('installed') or isinstance(exc, Cancelled)):
                if not added_registration: record['registration_owned'] = bool(previous.get('registration_owned'))
                if not added_trust: record['trust_owned'] = bool(previous.get('trust_owned'))
                for added, function, argument, field in (
                    (added_registration, lambda path: self.settings.restore_registration(path, record.get('previous_registration')), manifest, 'registration_owned'),
                    (added_trust, self.settings.remove_trust, cert_path, 'trust_owned')):
                    if added:
                        try:
                            function(argument); record[field] = False
                            if field == 'registration_owned': record['previous_registration'] = None
                        except (OSError, ValueError): rollback_error = True
                if isinstance(exc, Cancelled):
                    record['installed'] = bool(previous.get('installed')) and not renewed
                    if not record['installed']: self.app.state.word_panels.clear()
                # Rollback may have revoked the trust just verified above.
                self.app.state.word_https_check = {}
                try: self._save(record)
                except OSError: rollback_error = True
            message = 'Настройка отменена. Можно повторить подключение или продолжить работу без панели.' if isinstance(exc, Cancelled) else str(exc) if isinstance(exc, ValueError) else 'Не удалось выполнить настройку. Проверьте доступ к данным приложения и повторите подключение.'
            if record and record.get('installed') and isinstance(exc, OSError):
                message = 'Подключение настроено, но учебный документ не удалось открыть. Нажмите «Открыть новый учебный документ в Word».'
            if rollback_error: message += ' Не удалось полностью отменить собственные изменения. Нажмите «Удалить подключение», чтобы повторить очистку.'
            self._phase('cancelled' if isinstance(exc, Cancelled) else 'error', message)

    def _remove(self):
        try:
            self._check_cancel()
            self._phase('removing', 'Удаляем только созданное приложением подключение… Подтвердите запрос системы, если он появился.')
            record = self._record()
            if record:
                config, cert = word_setup.configuration(self.root, getattr(self.app.state, 'http_port', 8765))
                if record['fingerprint'] != cert.fingerprint(word_setup.hashes.SHA256()).hex(): raise ValueError('Сертификат изменился. Прежнее доверие и подключение сохранены; автоматическое удаление недоступно.')
                if record.get('registration_owned'):
                    self.settings.restore_registration(self.root / 'BLExhibitManager.Word.xml', record.get('previous_registration'))
                    record['installed'] = False
                    record['registration_owned'] = False; record['previous_registration'] = None; self._save(record)
                    self.app.state.word_panels.clear()
                if record.get('trust_owned'):
                    record['installed'] = False; self._save(record)
                    self.app.state.word_https_check = {}
                    self.settings.remove_trust(Path(config['cert']))
                    record['trust_owned'] = False; self._save(record)
                record['installed'] = False; self._save(record)
            self.app.state.word_panels.clear()
            self.app.state.word_https_check = {}
            self._phase('removed', 'Созданное приложением подключение удалено. Проекты, документы и прежние настройки сохранены.')
        except Cancelled:
            self._phase('cancelled', 'Удаление отменено до изменения настроек.')
        except (OSError, ValueError) as exc:
            self._phase('error', str(exc) if isinstance(exc, ValueError) else 'Не удалось удалить собственное подключение. Можно повторить удаление; документы сохранены.')

    def open(self):
        with self.lock:
            if self.thread and self.thread.is_alive(): raise ValueError('Дождитесь окончания настройки.')
            if not self.status().get('installed'): raise ValueError('Сначала подключите или обновите панель в настройках.')
            self.settings.open_word(starter_document(self.root))
            self._phase('waiting_word', 'Учебный документ открыт. Дождитесь панели в Word; статус обновится автоматически.')
            return self.status()


def manager(app, root):
    with word_setup.SETUP_LOCK:
        installer = getattr(app.state, 'word_installer', None)
        if not installer or installer.root != Path(root):
            installer = Installer(root, app); app.state.word_installer = installer
        return installer
