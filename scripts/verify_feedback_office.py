"""Native Office acceptance with fictional material; no converter substitute."""
import argparse
import json
import sys
import subprocess
import ctypes
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
from lxml import etree as E
from PIL import Image
from pypdf import PdfReader
import pypdfium2 as pdfium
import pypdfium2.raw as pdfium_raw
from exhibit import word, pdf, main_pdf
from exhibit.project import Store
from exhibit.samples import make_pdf
from tests.test_feedback import citation_docx

TEXTS=[
    'Annex 1, First Training Response, para. 1.',
    'RLA-12, Training Statute, Articles 1(1), 2(1);',
    'First Expert Report of Alex Sample, paras. 32 and 38.',
    'Amendment to the Training Response, paras. 5-6; R-19, Response to the Training Request, dated 10 April 2026, paras. 9-10.',
    'R-1, Training Statement, Annex No. 6 to Building Contract 100001, Clause 5; R-2, Training Statement, Annex No. 6 to Building Contract 100002, Clause 5.',
]
LABELS=['Annex 1','RLA-12','First Expert Report of Alex Sample','Amendment to the Training Response','R-19','R-1','R-2']
EXPECTED=['Annex 1, First Training Response','RLA-12, Training Statute',LABELS[2],LABELS[3],
    'R-19, Response to the Training Request, dated 10 April 2026',
    'R-1, Training Statement, Annex No. 6 to Building Contract 100001',
    'R-2, Training Statement, Annex No. 6 to Building Contract 100002']


def fixture():
    parts=word.package(citation_docx(TEXTS));root,rows=word.paragraphs(parts)
    for index,(*_,para) in enumerate(rows):
        # Split the identifier/title across mixed bold and italic runs.
        text=word.text_of(para)
        for child in list(para):para.remove(child)
        r=E.SubElement(para,f'{{{word.W}}}r');props=E.SubElement(r,f'{{{word.W}}}rPr')
        E.SubElement(props,f'{{{word.W}}}vertAlign',{f'{{{word.W}}}val':'superscript'})
        E.SubElement(r,f'{{{word.W}}}footnoteRef')
        r=E.SubElement(para,f'{{{word.W}}}r');E.SubElement(r,f'{{{word.W}}}t',{'{http://www.w3.org/XML/1998/namespace}space':'preserve'}).text=' '
        for i,value in enumerate((text[:text.index(',')],text[text.index(','):])):
            r=E.SubElement(para,f'{{{word.W}}}r');props=E.SubElement(r,f'{{{word.W}}}rPr')
            E.SubElement(props,f'{{{word.W}}}rFonts',{f'{{{word.W}}}ascii':'Times New Roman',f'{{{word.W}}}hAnsi':'Times New Roman'})
            E.SubElement(props,f'{{{word.W}}}'+('b' if i==0 else 'i'))
            E.SubElement(props,f'{{{word.W}}}sz',{f'{{{word.W}}}val':'20'})
            E.SubElement(r,f'{{{word.W}}}t',{'{http://www.w3.org/XML/1998/namespace}space':'preserve'}).text=value
    para=rows[0][-1]
    r=E.SubElement(para,f'{{{word.W}}}r');E.SubElement(r,f'{{{word.W}}}t',{'{http://www.w3.org/XML/1998/namespace}space':'preserve'}).text=' Website: '
    h=E.SubElement(para,f'{{{word.W}}}hyperlink',{f'{{{word.R}}}id':'rIdWebsite'})
    r=E.SubElement(h,f'{{{word.W}}}r');props=E.SubElement(r,f'{{{word.W}}}rPr')
    E.SubElement(props,f'{{{word.W}}}color',{f'{{{word.W}}}val':'0000FF'})
    E.SubElement(props,f'{{{word.W}}}u',{f'{{{word.W}}}val':'single'})
    E.SubElement(r,f'{{{word.W}}}t').text='example.org'
    rels=word.xml(parts[word.RELS]) if word.RELS in parts else E.Element(f'{{{word.REL}}}Relationships',nsmap={None:word.REL})
    E.SubElement(rels,f'{{{word.REL}}}Relationship',Id='rIdWebsite',Type=word.R+'/hyperlink',Target='https://example.org/',TargetMode='External')
    parts[word.FOOT]=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    parts[word.RELS]=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
    out=BytesIO()
    with ZipFile(out,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    return out.getvalue()


def create_review(store):
    p=store.create('Проверка замечаний коллеги')
    for label in LABELS:
        store.upload(p,label+'.pdf',make_pdf(label,[['Fictional material for feedback review.']]))
        d=p['documents'][-1]
        values={'title':label,'filename':label+'.pdf','folder':'Exhibits'}
        if label=='Annex 1':values.update(prefix='',number=1,designation='Annex')
        elif '-' in label:
            prefix,number=label.rsplit('-',1);values.update(prefix=prefix,number=int(number))
        else:values.update(prefix='',number=None,designation='')
        store.update(p,d['id'],values);store.approve(p,d['id'],'document')
    store.upload(p,'Main document.docx',fixture(),'main');store.scan(p)
    assert [r['link_text'] for r in p['references']]==EXPECTED
    store.confirm_links(p)
    return p


def verify(output,engine=None):
    output.mkdir(parents=True,exist_ok=True)
    actual=engine or main_pdf.available()
    assert actual['available'],actual
    (output/'Original Main document.docx').write_bytes(fixture())
    real_popen=main_pdf.popen
    with (output/'office.log').open('w',encoding='utf-8') as log,TemporaryDirectory() as folder,patch.object(main_pdf,'available',return_value=actual):
        store=Store(folder);project=create_review(store)
        def logged_popen(command,**kwargs):
            log.write(repr(command)+'\n');log.flush()
            kwargs.update(stdout=log,stderr=subprocess.STDOUT)
            return real_popen(command,**kwargs)
        with patch.object(main_pdf,'popen',side_effect=logged_popen):
            try:package=store.export(project)
            except Exception:
                log.flush();print((output/'office.log').read_text('utf-8',errors='replace'))
                if sys.platform=='darwin':
                    executable=Path(actual['executable'])
                    if executable.stat().st_size<10000:print('Launcher:',executable.read_text(errors='replace'))
                    direct='/Applications/LibreOffice.app/Contents/MacOS/soffice'
                    from docx import Document
                    simple=BytesIO();d=Document();d.add_paragraph('Fictional smoke check');d.save(simple)
                    linked,paths=store.linked_main(project)
                    marked,_=main_pdf.marked_docx(linked,list(paths.values()))
                    from exhibit.samples import main_docx
                    for name,data in [('simple',simple.getvalue()),('base',main_docx()),('citations',citation_docx(TEXTS)),('original',fixture()),('linked',linked),('marked',marked)]:
                        source=output/(name+'.docx');source.write_bytes(data)
                        profile=output/(name+'-profile')
                        result=subprocess.run([direct,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(output),str(source)],capture_output=True,timeout=45)
                        print(name,result.returncode,result.stdout.decode(errors='replace'),result.stderr.decode(errors='replace'),(output/(name+'.pdf')).exists())
                raise
        (output/'Feedback review.zip').write_bytes(package)
        with ZipFile(BytesIO(package)) as z:z.extractall(output)
        (output/'Original Main document.docx').write_bytes(fixture())
        (output/'Original Annex 1.pdf').write_bytes(store.source(project,project['documents'][0]['original']))
    linked=word.package((output/'Submission/Main document.docx').read_bytes())
    root=word.xml(linked[word.FOOT]);managed=[h for h in root.findall('.//w:hyperlink',word.NS) if h.get(f'{{{word.R}}}id')!='rIdWebsite']
    assert [word.text_of(h) for h in managed]==EXPECTED
    for h in managed:
        for run in h.findall('w:r',word.NS):
            props=run.find('w:rPr',word.NS)
            assert props.find('w:color',word.NS).attrib=={f'{{{word.W}}}val':'000000'}
            assert props.find('w:u',word.NS).get(f'{{{word.W}}}val')=='none'
            for name,value in [('b','1'),('bCs','1'),('i','0'),('iCs','0')]:
                assert props.find('w:'+name,word.NS).get(f'{{{word.W}}}val')==value
    old=word.xml(word.package(fixture())[word.FOOT]).find('.//w:hyperlink',word.NS)
    website=root.xpath('.//w:hyperlink[@r:id="rIdWebsite"]',namespaces=word.NS)[0]
    assert E.tostring(old)==E.tostring(website)
    data=(output/'Submission/Main document.pdf').read_bytes();reader=PdfReader(BytesIO(data));count=0;targets=set();websites=0;bold_characters=0
    for index,page in enumerate(reader.pages):
        png=pdf.render_png(data,index+1,2);(output/f'Main-page-{index+1}.png').write_bytes(png)
        rendered=Image.open(BytesIO(png)).convert('RGB')
        sx=rendered.width/float(page.mediabox.width);sy=rendered.height/float(page.mediabox.height)
        for item in page.get('/Annots',[]):
            a=item.get_object();action=a.get('/A',{})
            if action.get('/URI','').startswith('https://example.org'):websites+=1
            if action.get('/S')!='/GoToR':continue
            count+=1;target=action['/F']['/UF'];targets.add(target)
            assert (output/'Submission'/target).is_file()
            assert list(a['/Border'])==[0,0,0] and '/BS' not in a and '/C' not in a
            x0,y0,x1,y1=map(float,a['/Rect']);h=float(page.mediabox.height)
            pixels=rendered.crop((int(x0*sx),int((h-y1)*sy),int(x1*sx),int((h-y0)*sy)))
            assert not any(max(pixel)-min(pixel)>8 for pixel in pixels.get_flattened_data()),'Coloured managed link'
        rectangles=[list(map(float,item.get_object()['/Rect'])) for item in page.get('/Annots',[]) if item.get_object().get('/A',{}).get('/S')=='/GoToR']
        with pdfium.PdfDocument(data) as native:
            native_page=native[index];textpage=native_page.get_textpage()
            try:
                for ci in range(textpage.count_chars()):
                    if not chr(pdfium_raw.FPDFText_GetUnicode(textpage,ci)).isalnum():continue
                    x0,y0,x1,y1=textpage.get_charbox(ci);x,y=(x0+x1)/2,(y0+y1)/2
                    if not any(a<=x<=c and b<=y<=d for a,b,c,d in rectangles):continue
                    flags=ctypes.c_int();size=pdfium_raw.FPDFText_GetFontInfo(textpage,ci,None,0,ctypes.byref(flags));name=ctypes.create_string_buffer(size)
                    pdfium_raw.FPDFText_GetFontInfo(textpage,ci,name,size,ctypes.byref(flags))
                    weight=pdfium_raw.FPDFText_GetFontWeight(textpage,ci)
                    assert weight>=600 or b'bold' in name.value.lower(),('Link is not bold',name.value,weight)
                    assert not flags.value&64 and b'italic' not in name.value.lower(),('Link is italic',name.value)
                    bold_characters+=1
            finally:textpage.close();native_page.close()
    assert count>=len(EXPECTED) and len(targets)==len(LABELS) and websites==1,(count,targets,websites)
    assert bold_characters>100
    result={'engine':actual['label'],'citations':len(managed),'pdf_link_rectangles':count,'pdf_targets':len(targets),'pages':len(reader.pages),'black_links':True,'bold_upright_pdf_characters':bold_characters,'external_website_preserved':True}
    (output/'checks.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',default='output/feedback-review/word');args=parser.parse_args()
    verify(Path(args.output).resolve())
