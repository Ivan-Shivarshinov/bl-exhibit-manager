"""Synthetic layout investigation; no client files or text travel to CI."""
from io import BytesIO
from pathlib import Path
from hashlib import sha256
from zipfile import ZipFile, ZIP_DEFLATED
import argparse, json, subprocess, sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw
from lxml import etree as E
from exhibit import word, main_pdf, office_compat, font_diagnostics
from unittest.mock import patch
from pypdf import PdfReader

CT='http://schemas.openxmlformats.org/package/2006/content-types'

def pack(parts):
    out=BytesIO()
    with ZipFile(out,'w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
    return out.getvalue()

def footnotes(left=567, first=142, margin=72, long=False, font='Liberation Serif'):
    doc=Document();normal=doc.styles['Normal'];normal.font.name=font;normal.font.size=Pt(12)
    doc.sections[0].left_margin=Pt(margin);doc.sections[0].right_margin=Pt(72)
    for bold,italic in [(False,False),(True,False),(False,True),(True,True)]:
        p=doc.add_paragraph();r=p.add_run('Synthetic typography control');r.bold=bold;r.italic=italic
    p=doc.add_paragraph('Synthetic main text with a footnote')
    r=E.SubElement(p._p,f'{{{word.W}}}r');E.SubElement(r,f'{{{word.W}}}footnoteReference',{f'{{{word.W}}}id':'1'})
    root=E.Element(f'{{{word.W}}}footnotes',nsmap={'w':word.W})
    for id,kind,tag in [('-1','separator','separator'),('0','continuationSeparator','continuationSeparator')]:
        n=E.SubElement(root,f'{{{word.W}}}footnote',{f'{{{word.W}}}id':id,f'{{{word.W}}}type':kind})
        p=E.SubElement(n,f'{{{word.W}}}p');props=E.SubElement(p,f'{{{word.W}}}pPr')
        E.SubElement(props,f'{{{word.W}}}ind',{f'{{{word.W}}}left':str(left),f'{{{word.W}}}firstLine':str(first)})
        r=E.SubElement(p,f'{{{word.W}}}r');E.SubElement(r,f'{{{word.W}}}{tag}')
    n=E.SubElement(root,f'{{{word.W}}}footnote',{f'{{{word.W}}}id':'1'})
    for i in range(110 if long else 2):
        p=E.SubElement(n,f'{{{word.W}}}p');pp=E.SubElement(p,f'{{{word.W}}}pPr')
        E.SubElement(pp,f'{{{word.W}}}spacing',{f'{{{word.W}}}after':'0'})
        if i==0:
            r=E.SubElement(p,f'{{{word.W}}}r');E.SubElement(r,f'{{{word.W}}}footnoteRef')
        r=E.SubElement(p,f'{{{word.W}}}r');E.SubElement(r,f'{{{word.W}}}t').text=('Annex 1, Training source, p. 1. ' if i==0 else '')+f'Footnote evidence line {i+1:03}: synthetic material for layout verification.'
    buffer=BytesIO();doc.save(buffer);parts=word.package(buffer.getvalue());parts[word.FOOT]=E.tostring(root)
    types=word.xml(parts['[Content_Types].xml']);E.SubElement(types,'{'+CT+'}Override',PartName='/'+word.FOOT,ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    parts['[Content_Types].xml']=E.tostring(types)
    rels=word.xml(parts['word/_rels/document.xml.rels']);E.SubElement(rels,'{'+word.REL+'}Relationship',Id='rIdNotes',Type=word.R+'/footnotes',Target='footnotes.xml');parts['word/_rels/document.xml.rels']=E.tostring(rels)
    settings=word.xml(parts['word/settings.xml']);fp=E.SubElement(settings,f'{{{word.W}}}footnotePr')
    for i in ['-1','0']:E.SubElement(fp,f'{{{word.W}}}footnote',{f'{{{word.W}}}id':i})
    parts['word/settings.xml']=E.tostring(settings)
    return pack(parts)

def pdf_summary(data, directory=None):
    import pymupdf as fitz
    d=fitz.open(stream=data,filetype='pdf');r=PdfReader(BytesIO(data));pages=[]
    for i,p in enumerate(d):
        shapes=[]
        for drawing in p.get_drawings():
            shapes.append({'rect':list(drawing['rect']),'type':drawing['type'],'width':drawing['width'],'fill':drawing['fill'],'color':drawing['color']})
        pages.append({'content_sha256':sha256(r.pages[i].get_contents().get_data()).hexdigest(),
                      'box':list(r.pages[i].mediabox),'fonts':sorted(set(x[3] for x in p.get_fonts())),
                      'drawings':shapes,'images':[list(rect) for x in p.get_images() for rect in p.get_image_rects(x[0])],
                      'text_blocks':[[list(b[:4]),sha256(b[4].encode()).hexdigest()] for b in p.get_text('blocks') if isinstance(b[4],str)]})
        if directory:p.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(directory/f'page-{i+1}.png')
    return {'pages':len(d),'details':pages}

def signature(keep=False, font='Liberation Serif'):
    doc=Document();s=doc.sections[0];s.page_height=Pt(432);s.page_width=Pt(360)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Pt(36)
    normal=doc.styles['Normal'];normal.font.name=font;normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=Pt(12)
    for i in range(27):doc.add_paragraph(f'Synthetic body paragraph {i+1:03}.')
    label=doc.add_paragraph('Respectfully submitted,');label.paragraph_format.keep_with_next=keep
    p=doc.add_paragraph();p.paragraph_format.keep_with_next=keep
    img=Image.new('RGBA',(120,30),(255,255,255,0));ImageDraw.Draw(img).line([(5,25),(30,5),(20,25),(70,10),(50,25),(115,12)],fill=(20,40,100),width=3)
    b=BytesIO();img.save(b,format='PNG');b.seek(0)
    run=p.add_run();shape=run.add_picture(b,width=Pt(80));inline=shape._inline
    WP='http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
    anchor=E.Element('{'+WP+'}anchor',nsmap=inline.nsmap,simplePos='0',relativeHeight='0',behindDoc='0',locked='0',layoutInCell='1',allowOverlap='1',distT='0',distB='0',distL='0',distR='0')
    E.SubElement(anchor,'{'+WP+'}simplePos',x='0',y='0')
    for axis in ['H','V']:
        pos=E.SubElement(anchor,'{'+WP+'}position'+axis,relativeFrom='column' if axis=='H' else 'paragraph');E.SubElement(pos,'{'+WP+'}posOffset').text='0'
    for child in list(inline):anchor.append(child)
    anchor.insert(next(i for i,n in enumerate(anchor) if n.tag=='{'+WP+'}docPr'),E.Element('{'+WP+'}wrapNone'))
    inline.getparent().replace(inline,anchor)
    for text in ['____________________','Fictional Person','Counsel in synthetic example']:
        q=doc.add_paragraph(text);q.paragraph_format.keep_with_next=keep and text!='Counsel in synthetic example'
    out=BytesIO();doc.save(out);return out.getvalue()

def export(data, directory, engine, format='pdf'):
    directory.mkdir(parents=True,exist_ok=True)
    src=directory/'Main document.docx';src.write_bytes(data)
    dst=directory/'Main document.pdf'
    if format=='pdf':main_pdf.office_export(src,dst,directory,engine)
    else:
        profile=directory/'fodt-profile'
        subprocess.run([engine['executable'],'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','fodt','--outdir',str(directory),str(src)],check=True,timeout=120,stdout=subprocess.DEVNULL)
        return (directory/'Main document.fodt').read_bytes()
    return dst.read_bytes()

def main():
    import pymupdf as fitz
    ap=argparse.ArgumentParser();ap.add_argument('--output',required=True);ap.add_argument('--executable');args=ap.parse_args()
    out=Path(args.output).resolve();out.mkdir(parents=True,exist_ok=True)
    engine={'available':True,'engine':'libreoffice','label':'LibreOffice','executable':args.executable} if args.executable else main_pdf.available()
    report={'engine':engine,'version':subprocess.check_output([engine['executable'],'--version']).decode().strip() if engine['engine']=='libreoffice' else 'Word Windows','cases':{}}
    for name,left,first,margin,long in [('zero',0,0,72,False),('indent',567,142,72,False),('wide-margin',240,240,108,False),('continuation',567,142,72,True)]:
        data=footnotes(left,first,margin,long);dire=out/name
        pdf=export(data,dire,engine);report['cases'][name]=pdf_summary(pdf,dire)
        if engine['engine']=='libreoffice':(dire/'imported.fodt').write_bytes(export(data,dire/'import',engine,'fodt'))
        if engine['engine']=='libreoffice':
            with patch.object(main_pdf,'available',return_value=engine):fixed=main_pdf.convert_docx(data,[])
            dest=dire/'fixed';dest.mkdir(exist_ok=True);(dest/'Main document.pdf').write_bytes(fixed)
            got=pdf_summary(fixed,dest);report['cases'][name]['fixed']=got
            original=report['cases'][name]
            assert original['pages']==got['pages']
            for i,(old,new) in enumerate(zip(original['details'],got['details'])):
                assert old['text_blocks']==new['text_blocks'],(name,i,'text moved')
                assert old['fonts']==new['fonts'],(name,i,'font changed')
                expected=margin+(left+first)/20
                assert all(abs(d['rect'][0]-expected)<1 for d in new['drawings']),(name,i,new['drawings'])
                if i>0:assert all(abs(d['rect'][2]-540)<1 for d in new['drawings']),(name,i,'continuation')
            report['cases'][name]['unchanged_text_and_fonts']=True
    for keep in [False,True]:
        name='signature-'+str(keep).lower();dire=out/name
        report['cases'][name]=pdf_summary(export(signature(keep),dire,engine),dire)
        if keep:
            note=fitz.open(dire/'Main document.pdf');images=[i for i,p in enumerate(note) if any(p.get_image_rects(im[0]) for im in p.get_images())]
            for label in ['Respectfully submitted','Fictional Person','Counsel in synthetic example']:
                assert [i for i,p in enumerate(note) if label in p.get_text()]==images,('grouped signature split',label)
    for font in ['Times New Roman','Missing Synthetic Font']:
        name='font-'+font.replace(' ','-');dire=out/name;data=footnotes(font=font);result=export(data,dire,engine)
        report['cases'][name]={'preflight':font_diagnostics.inspect(data,engine),'rendered':font_diagnostics.pdf_report(result,[font]),'pdf':pdf_summary(result)}
        if font=='Missing Synthetic Font':
            assert report['cases'][name]['preflight']['missing']==[font]
            assert report['cases'][name]['rendered']['unconfirmed']==[font]
    from exhibit.project import Store
    from exhibit.samples import make_pdf
    from zipfile import ZipFile
    directory=out/'full-application';directory.mkdir(exist_ok=True)
    with patch.object(main_pdf,'available',return_value=engine):
        store=Store(directory/'data');p=store.create('Synthetic layout submission')
        attachment=make_pdf('Training source',[['Fictional evidence']]);store.upload(p,'Annex 1.pdf',attachment,mode='passthrough')
        data=footnotes();store.upload(p,'Main.docx',data,'main');store.scan(p);store.confirm_links(p)
        linked,paths=store.linked_main(p);direct=export(linked,directory/'direct-linked',engine)
        marked,markers=main_pdf.marked_docx(linked,list(paths.values()));prepared=office_compat.prepare_docx(marked) if engine['engine']=='libreoffice' else marked
        converted=export(prepared,directory/'prepared',engine);portable=main_pdf.portable_links(converted,markers)
        before,after=pdf_summary(converted),pdf_summary(portable)
        assert [x['content_sha256'] for x in before['details']]==[x['content_sha256'] for x in after['details']]
        assert max(abs(float(u)-float(v)) for x,y in zip(before['details'],after['details']) for u,v in zip(x['box'],y['box']))<1e-5
        archive=store.export(p);(directory/'Submission.zip').write_bytes(archive)
        with ZipFile(BytesIO(archive)) as z:
            downloaded=z.read('Submission/Main document.docx');pdf=z.read('Submission/Main document.pdf')
            assert z.read('Submission/Annex 1.pdf')==attachment
        original,actual=word.package(data),word.package(downloaded)
        assert all(actual[n]==b for n,b in original.items() if n not in (word.FOOT,word.RELS))
        assert store.source(p,p['main'])==data
        r=PdfReader(BytesIO(pdf));targets=[str(a.get_object()['/A']['/F']['/UF']) for page in r.pages for a in page.get('/Annots',[]) if a.get_object().get('/A',{}).get('/S')=='/GoToR']
        assert targets==['Annex 1.pdf'],targets
        (directory/'linked.docx').write_bytes(downloaded);(directory/'Main document.pdf').write_bytes(pdf)
        report['application']={'source_unchanged':True,'downloaded_docx_preserved':True,'portable_content_unchanged':True,'targets':targets,'direct_linked':pdf_summary(direct),'final':pdf_summary(pdf,directory),'fonts':store.main_pdf_status(p).get('rendered_fonts')}
    (out/'results.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    report['fonts']=font_diagnostics.inspect(footnotes(font='Times New Roman'),engine)
    (out/'results.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps({'version':report['version'],'cases':{k:v.get('pages',v.get('pdf',{}).get('pages')) for k,v in report['cases'].items()},'application':{k:v for k,v in report['application'].items() if k not in ('direct_linked','final')}}))

if __name__=='__main__':main()
