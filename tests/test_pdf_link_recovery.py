"""Synthetic tagged PDFs model the converter losing annotations, never note text."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import unittest

from docx import Document
from lxml import etree as E
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (ArrayObject, DictionaryObject, NameObject, NumberObject,
                           DecodedStreamObject)
from reportlab.pdfgen import canvas

from exhibit import main_pdf, word, pdf_link_recovery as recovery


def docx_fixture(texts=(('Annex 1', 'Annex 2'),), numbering=None):
    doc = Document()
    for i in range(len(texts)):
        run = doc.add_paragraph('Synthetic body').add_run()._r
        E.SubElement(run, '{'+word.W+'}footnoteReference', {'{'+word.W+'}id':str(i+1)})
    buffer = BytesIO(); doc.save(buffer); parts = word.package(buffer.getvalue())
    notes = E.Element('{'+word.W+'}footnotes', nsmap={'w':word.W, 'r':word.R})
    rels = E.Element('{'+word.REL+'}Relationships'); markers = {}
    for i, links in enumerate(texts):
        note = E.SubElement(notes, '{'+word.W+'}footnote', {'{'+word.W+'}id':str(i+1)})
        p = E.SubElement(note, '{'+word.W+'}p')
        E.SubElement(E.SubElement(p, '{'+word.W+'}r'), '{'+word.W+'}footnoteRef')
        for j, text in enumerate(links):
            marker = f'https://exhibit.invalid/{i}-{j}'; rid = f'r{i}-{j}'
            markers[marker] = f'Exhibits/Source {i}-{j}.pdf'
            E.SubElement(rels, '{'+word.REL+'}Relationship', Id=rid, Target=marker,
                         Type=word.R+'/hyperlink', TargetMode='External')
            link = E.SubElement(p, '{'+word.W+'}hyperlink', {'{'+word.R+'}id':rid})
            E.SubElement(E.SubElement(link, '{'+word.W+'}r'), '{'+word.W+'}t').text = text
            E.SubElement(E.SubElement(p, '{'+word.W+'}r'), '{'+word.W+'}t').text = '; '
    parts[word.FOOT] = E.tostring(notes); parts[word.RELS] = E.tostring(rels)
    if numbering:
        settings = word.xml(parts['word/settings.xml'])
        props = E.SubElement(settings, '{'+word.W+'}footnotePr')
        E.SubElement(props, '{'+word.W+'}numFmt', {'{'+word.W+'}val':numbering})
        parts['word/settings.xml'] = E.tostring(settings)
    out = BytesIO()
    with ZipFile(out, 'w', ZIP_DEFLATED) as z:
        for name, data in parts.items(): z.writestr(name, data)
    return out.getvalue(), markers


def pdf_fixture(notes=('1 Annex 1; Annex 2; ',), tagged=True, duplicate=False, existing=(), rotate=False):
    out = BytesIO(); c = canvas.Canvas(out, pagesize=(500, 700), invariant=1)
    c.drawString(40, 650, 'Body on preceding page'); c.showPage()
    for i, text in enumerate(notes):
        c._code.append(f'/Note << /MCID {i} >> BDC')
        c.drawString(40, 150-i*35, text)
        c._code.append('EMC')
    if duplicate: c.drawString(40, 400, notes[0])
    c.linkURL('https://example.org/preserved', (40, 20, 120, 30))
    for uri in existing: c.linkURL(uri, (40, 500, 100, 515))
    c.save(); writer = PdfWriter(clone_from=PdfReader(BytesIO(out.getvalue())))
    page = writer.pages[1]
    if rotate: page.rotate(90)
    if tagged:
        nodes = []
        for i in range(len(notes)):
            nodes.append(writer._add_object(DictionaryObject({
                NameObject('/S'):NameObject('/Note'), NameObject('/Pg'):page.indirect_reference,
                NameObject('/K'):NumberObject(i)})))
        root = writer._add_object(DictionaryObject({NameObject('/K'):ArrayObject(nodes)}))
        writer._root_object[NameObject('/StructTreeRoot')] = root
    out = BytesIO(); writer.write(out); return out.getvalue()


def remote_links(data):
    return [(pi, a.get_object()) for pi, page in enumerate(PdfReader(BytesIO(data)).pages)
            for a in page.get('/Annots', []) if a.get_object().get('/A', {}).get('/S') == '/GoToR']


class LinkRecovery(unittest.TestCase):
    def test_moved_note_multiple_links_content_and_existing_url_preserved(self):
        source, markers = docx_fixture(); raw = pdf_fixture()
        with self.assertRaisesRegex(ValueError, 'потеряны ссылки'): main_pdf.portable_links(raw, markers)
        fixed = main_pdf.portable_links(raw, markers, source)
        old = PdfReader(BytesIO(raw)); new = PdfReader(BytesIO(fixed))
        self.assertEqual(len(old.pages), len(new.pages))
        for a, b in zip(old.pages, new.pages):
            self.assertEqual(a.get_contents().get_data(), b.get_contents().get_data())
            self.assertEqual(list(a.mediabox), list(b.mediabox))
        links = remote_links(fixed); self.assertEqual(len(links), 2)
        self.assertEqual({a['/A']['/F']['/UF'] for _, a in links}, set(markers.values()))
        for pi, a in links:
            self.assertEqual(pi, 1)
            l, b, r, t = map(float, a['/Rect'])
            self.assertTrue(40 <= l < r < 200 and 140 < b < t < 165)
        self.assertEqual(new.pages[1]['/Annots'][0].get_object()['/A']['/URI'], 'https://example.org/preserved')

    def test_identical_note_text_disambiguated_by_number(self):
        source, markers = docx_fixture((('Annex 1',), ('Annex 1',)))
        raw = pdf_fixture(('1 Annex 1; ', '2 Annex 1; '))
        links = remote_links(main_pdf.portable_links(raw, markers, source))
        self.assertEqual(len(links), 2)
        by_target = {a['/A']['/F']['/UF']: float(a['/Rect'][1]) for _, a in links}
        self.assertGreater(by_target['Exhibits/Source 0-0.pdf'], by_target['Exhibits/Source 1-0.pdf'])

    def test_ambiguous_changed_missing_untagged_or_rotated_note_blocks(self):
        source, markers = docx_fixture()
        cases = [pdf_fixture(tagged=False), pdf_fixture(duplicate=True),
                 pdf_fixture(('1 Annex 1; Annex 3; ',)), pdf_fixture(('2 Annex 1; Annex 2; ',)),
                 pdf_fixture(('1 Annex 1; Annex 2; ', '1 Annex 1; Annex 2; ')),
                 pdf_fixture(rotate=True)]
        for raw in cases:
            with self.subTest(size=len(raw)), self.assertRaisesRegex(ValueError, 'Сноски: 1'):
                main_pdf.portable_links(raw, markers, source)

    def test_custom_numbering_blocks_only_recovery(self):
        source, markers = docx_fixture(numbering='lowerRoman')
        with self.assertRaisesRegex(ValueError, 'потеряны ссылки'):
            main_pdf.portable_links(pdf_fixture(), markers, source)
        fixed = main_pdf.portable_links(pdf_fixture(existing=list(markers)), markers, source)
        self.assertEqual(len(remote_links(fixed)), 2)

    def test_existing_markers_are_not_recovered_again(self):
        source, markers = docx_fixture()
        raw = pdf_fixture(existing=list(markers))
        self.assertEqual(recovery.recover(raw, markers, source), raw)
        self.assertEqual(len(remote_links(main_pdf.portable_links(raw, markers, source))), 2)

    def test_partial_recovery_still_blocks_the_whole_pdf(self):
        source, markers = docx_fixture((('Annex 1',), ('Annex 2',)))
        with self.assertRaisesRegex(ValueError, 'Сноски: 2'):
            main_pdf.portable_links(pdf_fixture(('1 Annex 1; ',)), markers, source)

    def test_conflicting_existing_link_is_not_covered(self):
        source, markers = docx_fixture()
        writer = PdfWriter(clone_from=PdfReader(BytesIO(pdf_fixture())))
        from pypdf.annotations import Link
        writer.add_annotation(1, Link(rect=(40, 140, 200, 170), url='https://example.org/other'))
        out = BytesIO(); writer.write(out)
        with self.assertRaisesRegex(ValueError, 'потеряны ссылки'):
            main_pdf.portable_links(out.getvalue(), markers, source)

    def test_misplaced_neighbour_moves_only_with_exact_proof_of_its_own_text(self):
        source, markers = docx_fixture((('Annex 1',), ('Annex 2',)))
        writer = PdfWriter(clone_from=PdfReader(BytesIO(pdf_fixture(('1 Annex 1; ', '2 Annex 2; ')))))
        from pypdf.annotations import Link
        # Converter assigned note 2's target to note 1's position.
        writer.add_annotation(1, Link(rect=(40, 140, 200, 170), url=list(markers)[1]))
        out = BytesIO(); writer.write(out)
        links = remote_links(main_pdf.portable_links(out.getvalue(), markers, source))
        self.assertEqual(len(links), 2)
        ys = {a['/A']['/F']['/UF']:float(a['/Rect'][1]) for _, a in links}
        self.assertGreater(ys['Exhibits/Source 0-0.pdf'], 140)
        self.assertLess(ys['Exhibits/Source 1-0.pdf'], 120)
        # Without the neighbour's complete text, neither link can be safely repaired.
        wrong = PdfWriter(clone_from=PdfReader(BytesIO(pdf_fixture(('1 Annex 1; ',)))))
        wrong.add_annotation(1, Link(rect=(40, 140, 200, 170), url=list(markers)[1]))
        out = BytesIO(); wrong.write(out)
        with self.assertRaisesRegex(ValueError, 'потеряны ссылки'):
            main_pdf.portable_links(out.getvalue(), markers, source)

    def test_multiline_link_has_separate_click_regions(self):
        source, markers = docx_fixture((('Annex 1 continuation',),))
        raw = pdf_fixture(('1 Annex 1 continuation; ',)); writer = PdfWriter(clone_from=PdfReader(BytesIO(raw)))
        stream = DecodedStreamObject()
        stream.set_data(b'/Note <</MCID 0>> BDC BT /F1 12 Tf 40 150 Td (1 Annex 1) Tj 0 -16 Td (continuation; ) Tj ET EMC')
        writer.pages[1][NameObject('/Contents')] = writer._add_object(stream)
        out = BytesIO(); writer.write(out)
        links = remote_links(main_pdf.portable_links(out.getvalue(), markers, source))
        self.assertEqual(len(links), 2)
        self.assertEqual(len({a['/A']['/F']['/UF'] for _, a in links}), 1)


if __name__ == '__main__': unittest.main()
