"""Fictional bitmap and TOC fixtures; never include client documents."""
from io import BytesIO
import struct
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches
from lxml import etree as E
from PIL import Image

from exhibit import word, main_pdf, office_compat


def plus_record(kind, flags=0, data=b''):
    return struct.pack('<HHII', kind, flags, len(data)+12, len(data)) + data


def bitmap_emf(image=None, extra=b'', pixel_format=0xE200B):
    if image is None:
        image = Image.new('RGBA', (120, 80), (225, 30, 20, 255))
        image.paste((20, 60, 230, 255), (60, 0, 120, 40))
        image.paste((30, 220, 70, 255), (0, 40, 60, 80))
        image.paste((230, 200, 20, 255), (60, 40, 120, 80))
    w, h = image.size
    rawmode = 'BGRa' if pixel_format == 0xE200B else 'BGRA'
    pixels = image.tobytes('raw', rawmode)
    bitmap = struct.pack('<IIiiiII', 0xDBC01002, 1, w, h, w*4, pixel_format, 0) + pixels
    commands = b''.join([
        plus_record(0x4001, 1, struct.pack('<4I', 0xDBC01002, 1, 96, 96)),
        plus_record(0x4030, 2, struct.pack('<f', 1)),
        plus_record(0x4009, 0, bytes([255, 255, 255, 0])),
        plus_record(0x4023),
        plus_record(0x4008, 0x800, struct.pack('<6I', 0xDBC01002, 0, 0, 0xFFFFFFFF, 0, 0)),
        plus_record(0x4008, 0x501, bitmap),
        extra,
        plus_record(0x401A, 1, struct.pack('<II8f', 0, 2, 0, 0, w, h, 0, 0, w, h)),
        plus_record(0x4002),
    ])
    comment = struct.pack('<III', 70, len(commands)+16, len(commands)+4) + b'EMF+' + commands
    end = struct.pack('<5I', 14, 20, 0, 0, 20)
    header = struct.pack('<II4i4i4IHH3I4i', 1, 88, 0, 0, w-1, h-1,
                         0, 0, round(w*2540/96), round(h*2540/96), 0x464D4520,
                         0x10000, 88+len(comment)+len(end), 3, 1, 0, 0, 0, 0,
                         1920, 1080, 508, 286)
    return header + comment + end


def package(parts):
    out = BytesIO()
    with ZipFile(out, 'w', ZIP_DEFLATED) as z:
        for name, content in parts.items(): z.writestr(name, content)
    return out.getvalue()


def fixture(emf=None):
    doc = Document()
    doc.add_heading('Illustrated training document', 0)
    p = doc.add_paragraph()
    def run(tag, text=None, **attrs):
        r = E.SubElement(p._p, f'{{{word.W}}}r')
        n = E.SubElement(r, f'{{{word.W}}}{tag}', {f'{{{word.W}}}{k}': v for k,v in attrs.items()})
        n.text = text
    run('fldChar', fldCharType='begin')
    run('instrText', ' TO')
    run('instrText', 'C \\h ')
    run('fldChar', fldCharType='separate')
    link = E.SubElement(p._p, f'{{{word.W}}}hyperlink', {f'{{{word.W}}}anchor': 'training'})
    r = E.SubElement(link, f'{{{word.W}}}r'); props = E.SubElement(r, f'{{{word.W}}}rPr')
    E.SubElement(props, f'{{{word.W}}}b')
    E.SubElement(props, f'{{{word.W}}}color', {f'{{{word.W}}}val': '467886', f'{{{word.W}}}themeColor': 'hyperlink'})
    E.SubElement(props, f'{{{word.W}}}u', {f'{{{word.W}}}val': 'single'})
    E.SubElement(r, f'{{{word.W}}}t').text = 'Training chapter'
    run('fldChar', fldCharType='end')
    p = doc.add_paragraph('An existing website: ')
    h = E.SubElement(p._p, f'{{{word.W}}}hyperlink', {f'{{{word.W}}}anchor': 'training'})
    r = E.SubElement(h, f'{{{word.W}}}r'); props = E.SubElement(r, f'{{{word.W}}}rPr')
    E.SubElement(props, f'{{{word.W}}}color', {f'{{{word.W}}}val': '0000FF'})
    E.SubElement(props, f'{{{word.W}}}u', {f'{{{word.W}}}val': 'single'})
    E.SubElement(r, f'{{{word.W}}}t').text = 'Keep this link blue'
    E.SubElement(p._p, f'{{{word.W}}}bookmarkStart', {f'{{{word.W}}}id':'1', f'{{{word.W}}}name':'training'})
    E.SubElement(p._p, f'{{{word.W}}}bookmarkEnd', {f'{{{word.W}}}id':'1'})
    png = BytesIO(); Image.new('RGB', (120,80), 'red').save(png, format='PNG'); png.seek(0)
    doc.add_picture(png, width=Inches(3))
    p = doc.add_paragraph('Training citation')
    r = E.SubElement(p._p, f'{{{word.W}}}r')
    E.SubElement(r, f'{{{word.W}}}footnoteReference', {f'{{{word.W}}}id':'1'})
    out = BytesIO(); doc.save(out)
    parts = word.package(out.getvalue())
    parts['word/media/image1.emf'] = bitmap_emf() if emf is None else emf
    del parts['word/media/image1.png']
    parts['word/_rels/document.xml.rels'] = parts['word/_rels/document.xml.rels'].replace(b'media/image1.png', b'media/image1.emf')
    types = word.xml(parts['[Content_Types].xml'])
    E.SubElement(types, '{'+office_compat.CT+'}Default', Extension='emf', ContentType='image/x-emf')
    from tests.test_feedback import citation_docx
    notes = word.package(citation_docx(['Annex 1, Training source, p. 1.']))
    parts[word.FOOT] = notes[word.FOOT]
    E.SubElement(types, '{'+office_compat.CT+'}Override', PartName='/'+word.FOOT,
                 ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
    rels = word.xml(parts['word/_rels/document.xml.rels'])
    E.SubElement(rels, '{'+word.REL+'}Relationship', Id='rIdTrainingNotes', Type=word.R+'/footnotes', Target='footnotes.xml')
    parts['word/_rels/document.xml.rels'] = E.tostring(rels)
    parts['[Content_Types].xml'] = E.tostring(types)
    return package(parts)


class OfficeCompatibility(unittest.TestCase):
    def test_bitmap_colors_orientation_and_alpha(self):
        image = Image.open(BytesIO(office_compat.emf_bitmap_png(bitmap_emf())))
        self.assertEqual(image.size, (120,80))
        self.assertEqual(image.getpixel((90,10)), (20,60,230,255))
        self.assertEqual(image.getpixel((10,60)), (30,220,70,255))
        for fmt in [0xE200B, 0x26200A]:
            translucent = Image.new('RGBA', (2,2), (255,0,0,128))
            result = Image.open(BytesIO(office_compat.emf_bitmap_png(bitmap_emf(translucent, pixel_format=fmt))))
            self.assertEqual(result.getpixel((0,0)), (255,0,0,128))

    def test_only_disposable_parts_change(self):
        source = fixture(); original = word.package(source)
        result = word.package(office_compat.prepare_docx(source))
        self.assertEqual(original['word/media/image1.emf'], result['word/media/image1.emf'])
        for name in original:
            if name not in ['word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml']:
                self.assertEqual(original[name], result[name], name)
        root = word.xml(result['word/document.xml'])
        links = root.findall('.//w:hyperlink', word.NS)
        self.assertEqual(links[0].find('.//w:color',word.NS).attrib, {f'{{{word.W}}}val':'000000'})
        self.assertEqual(links[0].find('.//w:u',word.NS).get(f'{{{word.W}}}val'), 'none')
        self.assertIsNotNone(links[0].find('.//w:b',word.NS))
        before = word.xml(original['word/document.xml']).findall('.//w:hyperlink',word.NS)
        self.assertEqual(E.tostring(links[1]), E.tostring(before[1]))
        self.assertEqual(E.tostring(root.find('.//w:drawing',word.NS)), E.tostring(word.xml(original['word/document.xml']).find('.//w:drawing',word.NS)))
        self.assertIn(b'media/image1.emf.png', result['word/_rels/document.xml.rels'])
        self.assertEqual(word.package(source), original)

    def test_unsupported_drawing_is_not_silently_replaced_by_its_bitmap(self):
        for extra in [plus_record(0x402A, 0, struct.pack('<2f',1,2)), plus_record(0x400A), plus_record(0x4004)]:
            with self.assertRaisesRegex(ValueError, 'PNG/JPEG'):
                office_compat.prepare_docx(fixture(bitmap_emf(extra=extra)))
        for data in [bitmap_emf()[:-4], bitmap_emf()[:100], b'not EMF']:
            with self.assertRaises(ValueError): office_compat.emf_bitmap_png(data)

    def test_ordinary_emf_and_noop_docx_are_preserved(self):
        data = bitmap_emf(); ordinary = data[:88] + struct.pack('<5I',14,20,0,0,20)
        self.assertIsNone(office_compat.emf_bitmap_png(ordinary))
        doc = Document(); doc.add_paragraph('Plain text'); out = BytesIO(); doc.save(out)
        self.assertEqual(out.getvalue(), office_compat.prepare_docx(out.getvalue()))

    def test_native_word_path_is_untouched_and_libreoffice_gets_copy(self):
        source = fixture()
        from reportlab.pdfgen import canvas
        def export(path, target, directory, engine):
            parts = word.package(path.read_bytes())
            self.assertEqual('word/media/image1.emf.png' in parts, engine['engine']=='libreoffice')
            out = BytesIO(); c = canvas.Canvas(out); c.drawString(10,10,'fixture'); c.save(); target.write_bytes(out.getvalue())
        for engine in ['word','libreoffice']:
            with patch.object(main_pdf,'available',return_value={'available':True,'engine':engine}), patch.object(main_pdf,'office_export',side_effect=export):
                self.assertTrue(main_pdf.convert_docx(source,[]).startswith(b'%PDF'))

    def test_conversion_version_invalidates_old_cached_pdf(self):
        from exhibit.project import Store
        from exhibit.samples import make_pdf
        with tempfile.TemporaryDirectory() as temp:
            store = Store(temp); project = store.create('Compatibility')
            store.upload(project,'Annex 1.pdf',make_pdf('Training',[['Sample']]),mode='passthrough')
            store.upload(project,'Main.docx',fixture(),'main'); store.scan(project); store.confirm_links(project)
            current = store.main_pdf_key(project)[0]
            with patch.object(main_pdf, 'CONVERSION_VERSION', 2):
                self.assertNotEqual(current, store.main_pdf_key(project)[0])

    def test_simple_toc_and_repeated_media_in_header(self):
        parts = word.package(fixture())
        simple = E.Element(f'{{{word.W}}}hdr', nsmap=word.NS)
        field = E.SubElement(simple,f'{{{word.W}}}fldSimple',{f'{{{word.W}}}instr':'TOC \\h'})
        r = E.SubElement(field,f'{{{word.W}}}r'); E.SubElement(r,f'{{{word.W}}}t').text='Header TOC'
        parts['word/header1.xml'] = E.tostring(simple)
        parts['word/_rels/header1.xml.rels'] = parts['word/_rels/document.xml.rels']
        result = word.package(office_compat.prepare_docx(package(parts)))
        self.assertEqual(sum(n.endswith('.emf.png') for n in result),1)
        self.assertIn(b'media/image1.emf.png', result['word/_rels/header1.xml.rels'])
        self.assertIsNotNone(word.xml(result['word/header1.xml']).find('.//w:color',word.NS))


if __name__ == '__main__': unittest.main()
