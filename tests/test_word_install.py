from io import BytesIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import threading
from unittest import TestCase
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from docx import Document
from exhibit import word_setup, word_tls
from exhibit.word_connection import WordPanels
from exhibit.word_install import Installer, NativeSettings, Cancelled, starter_document


class Settings:
    platform = 'test'
    def __init__(self):
        self.has_trust = self.has_registration = self.has_certificate = False
        self.calls = []
        self.failure = None
        self.waiting = threading.Event()

    def check_word(self):
        if self.failure == 'missing': raise ValueError('Word не найден.')
    def registration(self, manifest):
        if self.failure == 'conflict': raise ValueError('Другая регистрация сохранена.')
        return self.has_registration
    def trusted(self, cert, cancel=None): return self.has_trust
    def certificate_present(self, cert): return self.has_certificate
    def install_trust(self, cert, cancel):
        self.calls.append('trust'); self.has_trust = self.has_certificate = True
        if self.failure == 'cancel': raise Cancelled()
        if self.failure == 'waiting':
            self.waiting.set(); cancel.wait(3); raise Cancelled()
    def register(self, manifest):
        self.calls.append('register'); self.has_registration = True
        if self.failure == 'partial': raise OSError('Simulated denied registration')
    def unregister(self, manifest): self.calls.append('unregister'); self.has_registration = False
    def remove_trust(self, cert): self.calls.append('untrust'); self.has_trust = False
    def open_word(self, path):
        self.calls.append('open'); assert Document(path).paragraphs
        if self.failure == 'cancel_open': raise Cancelled()
        if self.failure == 'open': raise OSError('Simulated failed launch')


class WordInstallTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.app = FastAPI(); self.app.state.http_port = 8765
        self.app.state.word_panels = WordPanels()
        @self.app.get('/api/health')
        def health(): return {'application':'bl-exhibit-manager'}
        self.settings = Settings()
        self.installer = Installer(self.root, self.app, self.settings)
        self.app.state.word_installer = self.installer
        def context():
            import ssl
            return ssl.create_default_context(cafile=word_setup.configuration(self.root)[0]['cert'])
        self.trust = patch.object(word_setup, 'trust_context', side_effect=context)
        self.trust.start(); self.addCleanup(self.trust.stop)
        self.native_trust = patch.object(NativeSettings, 'trusted', return_value=True)
        self.native_trust.start(); self.addCleanup(self.native_trust.stop)
        self.addCleanup(self.stop_listener)

    def stop_listener(self):
        self.installer.cancel.set()
        if self.installer.thread: self.installer.thread.join(5)
        listener = getattr(self.app.state, 'word_listener', None)
        if listener: listener[0].should_exit = True; listener[1].join(5)

    def finish(self, remove=False):
        self.installer.begin(remove); self.installer.thread.join(10)
        self.assertFalse(self.installer.thread.is_alive(), 'Installer stuck')
        return self.installer.status()

    def test_install_repeat_update_and_remove_only_owned_objects(self):
        self.assertEqual(self.finish()['phase'], 'waiting_word')
        original = {p:p.read_bytes() for p in self.root.rglob('*') if p.is_file() and p.suffix in ('.key','.crt','.xml')}
        self.assertEqual(self.settings.calls, ['trust','register','open'])
        self.assertEqual(self.finish()['phase'], 'waiting_word')
        self.assertEqual(self.settings.calls, ['trust','register','open','open'])
        self.assertEqual(original, {p:p.read_bytes() for p in original})
        restarted = Installer(self.root, self.app, self.settings)
        self.assertTrue(restarted.status()['installed'])
        self.app.state.word_https_check = {}
        self.assertEqual(word_setup.status(self.root, self.app)['state'], 'trusted')
        self.assertEqual(self.finish(remove=True)['phase'], 'removed')
        self.assertFalse(self.settings.has_trust); self.assertFalse(self.settings.has_registration)
        self.assertEqual(original, {p:p.read_bytes() for p in original})
        self.assertEqual(self.finish()['phase'], 'waiting_word')

    def test_existing_trust_and_registration_are_never_removed(self):
        self.settings.has_trust = self.settings.has_certificate = self.settings.has_registration = True
        self.finish(); self.finish(remove=True)
        self.assertTrue(self.settings.has_trust); self.assertTrue(self.settings.has_registration)
        self.assertEqual(self.settings.calls, ['open'])

    def test_denied_or_cancelled_setup_rolls_back_own_partial_changes_and_retries(self):
        for failure in ('cancel','partial'):
            self.settings.failure = failure
            result = self.finish()
            self.assertIn(result['phase'], ('cancelled','error'))
            self.assertFalse(self.settings.has_trust); self.assertFalse(self.settings.has_registration)
            self.assertNotIn('open', self.settings.calls)
            self.settings.failure = None
            self.assertEqual(self.finish()['phase'], 'waiting_word')
            self.finish(remove=True); self.settings.calls.clear()

    def test_missing_word_conflict_and_existing_trust_rule_leave_system_untouched(self):
        for failure in ('missing','conflict'):
            self.settings.failure = failure
            self.assertEqual(self.finish()['phase'], 'error')
            self.assertEqual(self.settings.calls, [])
        self.settings.failure = None; self.settings.has_certificate = True
        self.assertEqual(self.finish()['phase'], 'error')
        self.assertEqual(self.settings.calls, [])

    def test_cancel_does_not_wait_for_system_prompt_or_block_status(self):
        self.settings.failure = 'waiting'; self.installer.begin()
        self.assertTrue(self.settings.waiting.wait(5))
        self.assertEqual(word_setup.status(self.root, self.app)['installation']['phase'], 'system_confirmation')
        self.installer.stop(); self.installer.thread.join(5)
        self.assertEqual(self.installer.status()['phase'], 'cancelled')
        self.assertFalse(self.settings.has_trust)

    def test_failed_word_launch_keeps_successful_setup_for_explicit_retry(self):
        self.settings.failure = 'open'
        self.assertEqual(self.finish()['phase'], 'error')
        self.assertTrue(self.installer.status()['installed'])
        self.assertTrue(self.settings.has_trust); self.assertTrue(self.settings.has_registration)
        self.settings.failure = None; self.installer.open()
        self.assertEqual(self.installer.status()['phase'], 'waiting_word')
        self.assertEqual(self.settings.calls.count('trust'), 1)

    def test_corrupt_ownership_marker_never_changes_system(self):
        (self.root/'word-installation.json').write_text('{"platform":"test","trust_owned":"yes"}', 'utf-8')
        self.assertEqual(self.finish()['phase'], 'error')
        self.assertEqual(self.settings.calls, [])

    def test_cancellation_at_final_step_rolls_back_a_new_installation(self):
        self.settings.failure = 'cancel_open'
        self.assertEqual(self.finish()['phase'], 'cancelled')
        self.assertFalse(self.installer.status()['installed'])
        self.assertFalse(self.settings.has_trust)
        self.assertFalse(self.settings.has_registration)

    def test_starter_documents_are_new_valid_docx_with_visible_registered_panel(self):
        first = starter_document(self.root); first.write_bytes(first.read_bytes()+b'USER_EDIT')
        before = first.read_bytes(); second = starter_document(self.root)
        self.assertNotEqual(first, second); self.assertEqual(first.read_bytes(), before)
        self.assertIn('учебный', Document(second).paragraphs[1].text)
        with ZipFile(second) as archive:
            self.assertIn(b'visibility="1"', archive.read('word/webextensions/taskpanes.xml'))
            self.assertIn(b'store="developer"', archive.read('word/webextensions/webextension.xml'))
            self.assertIn(b'webextensiontaskpanes', archive.read('_rels/.rels'))
            self.assertNotIn(b'PRIVATE KEY', second.read_bytes())

    def test_api_requires_explicit_consent_before_starting_setup(self):
        from exhibit.app import app, store
        with patch.object(store, 'root', self.root), patch.object(app.state, 'word_installer', self.installer, create=True):
            client = TestClient(app, base_url='http://localhost')
            for body in ({}, {'consent':False}, {'consent':1}, {'consent':True,'extra':True}):
                response = client.post('/api/setup/word/install', json=body, headers={'X-Exhibit-Local':'1'})
                self.assertEqual(response.status_code, 400)
            self.assertEqual(client.post('/api/setup/word/install', json={'consent':True}).status_code, 403)
            self.assertEqual(self.settings.calls, [])
            self.assertEqual(client.get('/api/health').status_code, 200)

    def test_mac_registration_preserves_other_files_and_rejects_conflict(self):
        settings = NativeSettings(platform='darwin', home=self.root/'home')
        manifest = self.root/'manifest.xml'; manifest.write_bytes(b'Synthetic manifest')
        target = settings._mac_target(); target.parent.mkdir(parents=True)
        foreign = target.parent/'foreign.xml'; foreign.write_bytes(b'Keep')
        settings.register(manifest); settings.register(manifest)
        self.assertEqual(target.read_bytes(), manifest.read_bytes())
        target.write_bytes(b'Different registration')
        with self.assertRaises(ValueError): settings.register(manifest)
        with self.assertRaises(ValueError): settings.unregister(manifest)
        self.assertEqual(target.read_bytes(), b'Different registration')
        target.write_bytes(manifest.read_bytes()); settings.unregister(manifest)
        self.assertFalse(target.exists()); self.assertEqual(foreign.read_bytes(), b'Keep')

    def test_mac_word_version_preflight(self):
        import plistlib
        settings = NativeSettings(platform='darwin', home=self.root)
        app = self.root/'Synthetic Word.app'; info = app/'Contents/Info.plist'
        info.parent.mkdir(parents=True)
        with patch.object(settings, 'word', return_value=app):
            for version, supported in [('16.69',False), ('16.70',True), ('16.101.3',True), ('invalid',False)]:
                info.write_bytes(plistlib.dumps({'CFBundleShortVersionString':version}))
                if supported: self.assertEqual(settings.check_word(), app)
                else:
                    with self.assertRaises(ValueError): settings.check_word()

    def test_mac_removal_disables_only_own_localhost_ssl_trust(self):
        settings = NativeSettings(platform='darwin', home=self.root)
        cert = self.root/'Synthetic.crt'
        with patch.object(settings, 'trusted', side_effect=[True,False]), patch.object(settings, 'run') as run:
            settings.remove_trust(cert)
            self.assertEqual(run.call_args.args[0], ['/usr/bin/security','add-trusted-cert','-r','deny','-p','ssl','-s','localhost','-k',str(settings.keychain),str(cert)])
        with patch.object(settings, 'trusted', return_value=True), patch.object(settings, 'run'):
            with self.assertRaises(ValueError): settings.remove_trust(cert)
